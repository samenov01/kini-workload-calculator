"""Новые МУПы и полный HTTP-путь загрузки/расчёта/скачивания без браузера."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from urllib.error import HTTPError
from urllib.parse import quote
from urllib.request import Request, urlopen
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter

from kini.model import read_plan
from webapp import Application, Handler, ThreadingHTTPServer, inspect_file

ROOT = Path(__file__).resolve().parent
FIXTURE_CONFIG = json.loads((ROOT / 'kini_config.json').read_text(encoding='utf-8'))
HAS_WORKBOOKS = all((ROOT / p).is_file() for p in
                    [FIXTURE_CONFIG['template'], *(g['file'] for g in FIXTURE_CONFIG['groups'])])


def new_plan(path, shift=0, extra_sheet=False):
    """Шесть семестров, сдвинутая шапка, новая дисциплина в строке 250, без кэшей Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Новый учебный план'
    ws.cell(3, 2, 'Computer Science · Год поступления: 2031')
    headers = {3:'Код дисциплины', 5:'Наименование дисциплины', 8:'ECTS', 9:'Экзамен',
               12:'ЛК', 13:'ПЗ', 14:'ЛЗ', **{17+i:f'{i+1} семестр' for i in range(6)}}
    for col, text in headers.items():
        ws.cell(25, col+shift, text)
    for sem in range(1, 7):
        r = 25+sem if sem < 6 else 250
        ws.cell(r, 3+shift, f'NEW.{sem}')
        ws.cell(r, 5+shift, f'Новая дисциплина {sem}')
        ws.cell(r, 9+shift, 'экзамен')
        ws.cell(r, 16+sem+shift, '3,5' if sem == 6 else 3)
        ws.cell(r, 8+shift, f'=SUM({get_column_letter(17+shift)}{r}:{get_column_letter(22+shift)}{r})')
        for col, value in [(12, 15), (13, 30), (14, 0)]:
            ws.cell(r, col+shift, value)
    if extra_sheet:
        wb.copy_worksheet(ws).title = 'Вторая программа'
    wb.save(path)
    wb.close()


class NewPlanTests(unittest.TestCase):
    def test_renamed_shifted_six_semester_plan_and_uncached_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Новый файл без известного имени.xlsx'
            for shift in (0, 3):
                with self.subTest(column_shift=shift):
                    new_plan(path, shift)
                    info = inspect_file(path)['sheets'][0]
                    self.assertEqual((info['program'], info['intake'], info['semesters']), ('CS', 2031, 6))
                    group = dict(name='CS-31', program='CS', intake=2031, students=8, language='KZ')
                    courses = read_plan(group, path, [])
                    self.assertEqual([c.semester for c in courses], [1, 2, 3, 4, 5, 6])
                    self.assertEqual([c.credits for c in courses], [3, 3, 3, 3, 3, 3.5])
                    self.assertEqual(courses[-1].row, 250)
                    self.assertTrue(all(c.hours == (15, 30, 0) for c in courses))

    def test_multiple_plans_require_explicit_sheet_and_unknown_header_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Две программы.xlsx'
            new_plan(path, extra_sheet=True)
            self.assertEqual(len(inspect_file(path)['sheets']), 2)
            group = dict(name='NEW-31', program='OTHER', intake=2031, students=8, language='RU')
            with self.assertRaisesRegex(ValueError, 'Выберите лист'):
                read_plan(group, path, [])
            group['sheet'] = 'Вторая программа'
            self.assertEqual(len(read_plan(group, path, [])), 6)
            wb = openpyxl.Workbook()
            wb.active['A1'] = 'Неизвестный формат'
            wb.save(path)
            wb.close()
            with self.assertRaisesRegex(ValueError, 'не распознан'):
                inspect_file(path)


@unittest.skipUnless(HAS_WORKBOOKS, 'Нужны локальные МУПы и Excel-шаблон; см. README.')
class BrowserAPITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        config = json.loads((ROOT / 'kini_config.json').read_text(encoding='utf-8'))
        config['template'] = str(ROOT / config['template'])
        for g in config['groups']:
            g['file'] = str(ROOT / g['file'])
        (root / 'kini_config.json').write_text(json.dumps(config), encoding='utf-8')
        cls.app = Application(root)
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        cls.server.app = cls.app
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = f'http://127.0.0.1:{cls.server.server_port}'
        cls.state = cls.request('/api/state')

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.app.executor.shutdown()
        cls.tmp.cleanup()

    @classmethod
    def request(cls, path, data=None, headers=None, raw=False):
        if data is not None and not isinstance(data, bytes):
            data = json.dumps(data).encode()
        h = {'X-KINI-Token': cls.app.token}
        h.update(headers or {})
        with urlopen(Request(cls.url + path, data=data, headers=h), timeout=30) as response:
            body = response.read()
            return body if raw else json.loads(body)

    def test_invalid_requests_are_rejected(self):
        for path, data, headers, code in [
            ('/api/generate', {'groups':self.state['groups']}, {'X-KINI-Token':'wrong'}, 403),
            ('/api/generate', [], {}, 400),
            ('/api/generate', {'groups':[None]}, {}, 400),
            ('/api/upload', b'not an excel file', {'X-Filename':'broken.xlsx'}, 400),
            ('/api/upload', b'test', {'X-Filename':'test.xls'}, 400),
            ('/api/state', None, {'Host':'external.example'}, 403),
        ]:
            with self.subTest(path=path, data=str(data)[:20]):
                with self.assertRaises(HTTPError) as caught:
                    self.request(path, data, headers)
                self.assertEqual(caught.exception.code, code)
                caught.exception.close()
        groups = deepcopy(self.state['groups'])
        groups[1]['name'] = groups[0]['name']
        with self.assertRaises(HTTPError) as caught:
            self.request('/api/generate', {'groups':groups})
        self.assertEqual(caught.exception.code, 400)
        caught.exception.close()

    def test_upload_generate_download_and_persist_selection(self):
        previous_revision = self.request('/api/state')['selection_revision']
        path = Path(self.tmp.name) / 'МУП нового набора.xlsx'
        new_plan(path, shift=2)
        file = self.request('/api/upload', path.read_bytes(), {'X-Filename':quote(path.name)})
        self.assertNotIn('path', file)
        group = dict(name='CS-31', program='CS', file_id=file['id'], sheet=file['sheets'][0]['sheet'],
                     intake=2031, students=11, language='KZ')
        job = self.request('/api/generate', {'groups':[group]})
        deadline = time.monotonic() + 60
        while job['status'] in ('queued', 'running') and time.monotonic() < deadline:
            time.sleep(.15)
            job = self.request('/api/jobs/' + job['id'])
        self.assertEqual(job['status'], 'complete', job['message'])
        self.assertEqual([r['year'] for r in job['results']], [2031, 2032, 2033])
        for result in job['results']:
            wb = openpyxl.load_workbook(BytesIO(self.request(result['url'], raw=True)), data_only=True)
            self.assertEqual(wb.sheetnames, ['Расчет'])
            if result['year'] == 2031:
                self.assertAlmostEqual(result['totals']['1'], 50.6)
                self.assertAlmostEqual(result['totals']['2'], 50.6)
                self.assertAlmostEqual(result['totals']['0'], 101.2)
                row = result['structure']['detail']['1'][0]
                self.assertEqual(wb['Расчет'].cell(row, 7).value, 11)
                self.assertEqual(wb['Расчет'].cell(row, 11).value, 4.5)
                self.assertEqual(wb['Расчет'].cell(row, 12).value, 1.1)
            wb.close()
        archive = zipfile.ZipFile(BytesIO(self.request(job['archive'], raw=True)))
        self.assertEqual(len([n for n in archive.namelist() if n.endswith('.xlsx')]), 3)
        self.assertIn('Замечания.json', archive.namelist())
        report = self.request(job['report'])
        self.assertEqual(len(report['courses']), 6)
        with self.assertRaises(HTTPError) as caught:
            self.request(f'/api/jobs/{job["id"]}/files/config.json')
        caught.exception.close()
        other = Application(Path(self.tmp.name))
        try:
            saved = other.state()
            self.assertEqual(saved['groups'][0]['students'], 11)
            self.assertNotEqual(saved['selection_revision'], previous_revision)
            self.assertTrue(any(f['id'] == file['id'] for f in saved['files']))
        finally:
            other.executor.shutdown()


if __name__ == '__main__':
    unittest.main(verbosity=2)
