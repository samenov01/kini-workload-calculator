"""Independent source inventory and comparison with the supplied KINI template."""
from collections import Counter
import ast
import json
import operator
from pathlib import Path
import re

import openpyxl
from openpyxl.formula import Tokenizer

from .layout import PlanValues, detect_layout, numeric
from .model import canonical, classify, cohorts, normal, number, tokens


def inventory_plan(path, group):
    """Read all subject rows, including shared credits and language components.

    This deliberately does not call read_plan or use its exclusion decisions.
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    fw = openpyxl.load_workbook(path)
    try:
        ws = wb[group['sheet']] if group.get('sheet') else next(
            s for s in wb if _recognized(s))
        layout = detect_layout(ws)
        resolver = PlanValues(wb, fw)
        merged = {(r, c): (m.min_row, m.min_col) for m in ws.merged_cells.ranges
                  for r in range(m.min_row, m.max_row + 1)
                  for c in range(m.min_col, m.max_col + 1)}

        def value(r, c):
            return resolver.get(ws.title, ws.cell(*merged.get((r, c), (r, c))).coordinate)

        records = []
        for r in range(layout.header + 1, ws.max_row + 1):
            name = str(ws.cell(r, layout.name).value or '').strip()
            code = str(ws.cell(r, layout.code).value or '').strip()
            if not name and not code:
                continue
            n = name.lower()
            structural = any(t in n for t in ('модул', 'module', 'цикл ', 'компонент', 'дополнительные виды обучения'))
            if structural and not re.search(r'\d{4}\b|[A-Za-zА-Яа-яІі]\d+\.\d+', code):
                continue
            if re.search(r'\bитого\b|барлығы|\btotal\b|келісілген|согласовано', n) or n in ('engineering', 'major'):
                continue
            semesters = {s: numeric(value(r, c)) for s, c in layout.semesters.items()
                         if numeric(value(r, c)) and numeric(value(r, c)) > 0}
            credits = numeric(value(r, layout.credits))
            if not semesters and not (code and number(credits)):
                continue
            records.append(dict(group=group['name'], program=group['program'], intake=group['intake'],
                file=path.name, sheet=ws.title, row=r, name=name, code=code, credits=credits,
                semesters=semesters, hours=[numeric(value(r, c)) for c in layout.hours],
                credit_anchor=merged.get((r, layout.credits), (r, layout.credits))[0]))
        return records
    finally:
        wb.close()
        fw.close()


def _recognized(sheet):
    try:
        detect_layout(sheet)
        return True
    except ValueError:
        return False


def verify_coverage(inventory, courses):
    """Every source subject/semester must be represented exactly once."""
    expected = {(r['group'], r['row'], s) for r in inventory for s in r['semesters']}
    actual = []
    for c in courses:
        actual.append((c.group, c.row, c.semester))
        actual.extend((c.group, s['row'], s['semester']) for s in c.shared_sources)
    counts = Counter(actual)
    return dict(expected=len(expected), represented=len(counts),
        missing=sorted(expected - set(counts)), extra=sorted(set(counts) - expected),
        duplicates=sorted(k for k, v in counts.items() if v != 1))


def compare_template(template, inventory, config):
    wb = openpyxl.load_workbook(template, data_only=True)
    try:
        ws = wb['Расчет']
        rows, term = [], None
        for r in range(13, ws.max_row + 1):
            heading = normal(ws.cell(r, 1).value).replace('і', 'i')
            if heading in ('iсеместр', 'iiсеместр'):
                term = 1 if heading == 'iсеместр' else 2
            group_set = cohorts(ws.cell(r, 5).value)
            if not group_set or not number(ws.cell(r, 3).value) or not ws.cell(r, 2).value:
                continue
            name = str(ws.cell(r, 2).value)
            key = canonical(name, config.get('aliases'))
            period = (int(ws.cell(r, 4).value) - 1) * 2 + term
            matches = []
            configured = Counter()
            for g in config['groups']:
                configured[g['program'], g['intake']] += g['students']
            expected_students = sum(configured[c] for c in group_set) if group_set <= configured.keys() else None
            for program, intake in sorted(group_set):
                pool = [i for i in inventory if i['program'] == program and i['intake'] == intake]
                same_intake = bool(pool)
                if not pool:
                    pool = [i for i in inventory if i['program'] == program]
                found = []
                for item in pool:
                    source_key = canonical(item['name'], config.get('aliases'))
                    verified = any(rule['program'] == program and rule['semester'] in item['semesters']
                        and normal(rule['name']) == normal(item['name']) and rule['key'] == key
                        for rule in config.get('reference_matches', []))
                    if tokens(name) & tokens(item['name']) or (key and key == source_key) or verified:
                        found.append(dict(source=item, semester_matches=period in item['semesters'],
                            credits_match=item['semesters'].get(period) == ws.cell(r, 3).value,
                            hours_match=(tuple(v or 0 for v in item['hours']) == tuple(ws.cell(r, col).value or 0 for col in (8, 9, 10)))
                                if len(item['semesters']) == 1 and any(v is not None for v in item['hours']) else None,
                            match='проверенное соответствие' if verified else 'совпадение названия/словаря'))
                combined = []
                if key in ('german1', 'german2') and ws.cell(r, 6).value == 'DE':
                    parts = [i for i in pool if canonical(i['name']) in ('german1', 'german2')
                             and period in i['semesters']]
                    if ({canonical(i['name']) for i in parts} == {'german1', 'german2'}
                            and len(parts) == 2 and sum(i['semesters'][period] for i in parts) == ws.cell(r, 3).value):
                        combined = parts
                        for item in parts:
                            if not any(f['source']['row'] == item['row'] for f in found):
                                found.append(dict(source=item, semester_matches=True, credits_match=False,
                                                  hours_match=None, match='часть исходного блока 4+2 кредита'))
                if not same_intake:
                    status = 'Нет МУПа этого набора'
                elif not found:
                    status = 'Точного соответствия нет'
                elif combined:
                    status = 'Совпадает (сумма частей)'
                elif any(f['semester_matches'] and f['credits_match'] for f in found):
                    status = 'Совпадает'
                elif any(f['semester_matches'] for f in found):
                    status = 'Отличаются кредиты'
                else:
                    status = 'Другой семестр'
                matches.append(dict(program=program, intake=intake, status=status,
                                    same_intake=same_intake, candidates=found,
                                    combined_credits=sum(i['semesters'][period] for i in combined) if combined else None))
            norms = []
            if classify(name) == 'ordinary':
                for column, expected in ((11, ws.cell(r, 3).value * 1.5), (12, ws.cell(r, 7).value * .1)):
                    actual = ws.cell(r, column).value
                    if not number(actual) or abs(actual - expected) > 1e-6:
                        norms.append(dict(cell=ws.cell(r, column).coordinate, actual=actual, expected=expected))
            rows.append(dict(row=r, name=name, term=term, semester=period,
                credits=ws.cell(r, 3).value, language=ws.cell(r, 6).value,
                students=ws.cell(r, 7).value, expected_students=expected_students,
                hours=[ws.cell(r, c).value for c in range(8, 20)], matches=matches, norm_differences=norms))
        return rows
    finally:
        wb.close()


def audit_formulas(path):
    """Check cached formula values against cached inputs and find dependency cycles.

    A stale cache is reported, never substituted as an approved teaching norm.
    """
    fw = openpyxl.load_workbook(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    graph, mismatches, unsupported = {}, [], []
    checked = 0
    ops = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
           ast.Div: operator.truediv, ast.Pow: operator.pow}

    def calculate(n):
        if isinstance(n, ast.Constant) and number(n.value):
            return n.value
        if isinstance(n, ast.BinOp) and type(n.op) in ops:
            return ops[type(n.op)](calculate(n.left), calculate(n.right))
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.UAdd, ast.USub)):
            return calculate(n.operand) * (-1 if isinstance(n.op, ast.USub) else 1)
        if isinstance(n, ast.Call) and isinstance(n.func, ast.Name):
            args = [calculate(a) for a in n.args]
            if n.func.id.upper() == 'SUM':
                return sum(args)
            if n.func.id.upper() == 'ROUND':
                return round(args[0], int(args[1]))
        raise ValueError('Unsupported formula')

    try:
        for ws in fw:
            for row in ws:
                for cell in row:
                    if cell.data_type != 'f':
                        continue
                    address = (ws.title, cell.coordinate)
                    dependencies, expr = set(), ''
                    try:
                        for token in Tokenizer(cell.value).items:
                            if token.type == 'OPERAND' and token.subtype == 'RANGE':
                                sheet, sep, ref = token.value.rpartition('!')
                                sheet = sheet.strip("'").replace("''", "'") if sep else ws.title
                                ref = ref if sep else token.value
                                if '[' in sheet:
                                    raise ValueError('External workbook')
                                cells = wb[sheet][ref.replace('$', '')]
                                cells = [cells] if not isinstance(cells, tuple) else [c for rr in cells for c in rr]
                                dependencies.update((sheet, c.coordinate) for c in cells)
                                expr += repr(sum(c.value for c in cells if number(c.value)))
                            else:
                                expr += token.value
                        graph[address] = dependencies
                        value = calculate(ast.parse(expr.replace('^', '**'), mode='eval').body)
                        cached = wb[ws.title][cell.coordinate].value
                        checked += 1
                        if not number(cached) or abs(cached - value) > 1e-6:
                            mismatches.append(dict(sheet=ws.title, cell=cell.coordinate, formula=cell.value,
                                                   cached=cached, from_cached_inputs=value))
                    except (ValueError, SyntaxError, KeyError, TypeError, ZeroDivisionError):
                        graph[address] = dependencies
                        unsupported.append(dict(sheet=ws.title, cell=cell.coordinate, formula=cell.value))
        cycles, visited, active, stack = set(), set(), set(), []

        def visit(node):
            if node in active:
                cycles.update(stack[stack.index(node):])
                return
            if node in visited:
                return
            active.add(node)
            stack.append(node)
            for dep in graph.get(node, ()):
                if dep in graph:
                    visit(dep)
            stack.pop()
            active.remove(node)
            visited.add(node)
        for node in graph:
            visit(node)
        return dict(checked=checked, total=len(graph), mismatches=mismatches,
                    unsupported=unsupported, cycles=sorted(cycles))
    finally:
        fw.close()
        wb.close()


def write_audit(directory, inventory, courses, template, config):
    directory = Path(directory)
    coverage = verify_coverage(inventory, courses)
    if any(coverage[k] for k in ('missing', 'extra', 'duplicates')):
        raise AssertionError('Неполное покрытие МУПа: ' + json.dumps(coverage, ensure_ascii=False))
    result = dict(coverage=coverage, template_rows=compare_template(template, inventory, config),
                  template_formulas=audit_formulas(template), inventory=inventory)
    (directory / 'Проверка МУПов и образца.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    def clean(value):
        return ' '.join(str(value).split()).replace('|', '/')
    f = result['template_formulas']
    lines = ['# Сверка МУПов с образцом KINI', '',
        f"Покрытие исходных дисциплин и семестров: {coverage['represented']} из {coverage['expected']}; потерь и повторов нет.", '',
        'Все составляющие общих блоков видны в Excel отдельными строками. Общие кредиты и часы начисляются один раз. Это не подтверждение выбора конкретной альтернативы группой.', '',
        f"Формулы образца: {f['checked']} проверено из {f['total']}; расхождений кэша с текущими сохранёнными входами: {len(f['mismatches'])}; ячеек в циклах: {len(f['cycles'])}; неподдержанных формул: {len(f['unsupported'])}.", '',
        'Проверка кэша не заменяет пересчёт книги в Excel и не утверждает нормативы. Группы другого набора не подменяются выбранными МУПами.', '',
        '| Строка образца | Предмет | Группа образца | Результат | Строки исходного МУПа |',
        '|---:|---|---|---|---|']
    for row in result['template_rows']:
        for match in row['matches']:
            sources = '; '.join(f"{c['source']['file']}:{c['source']['row']} (набор {c['source']['intake']}; семестры {','.join(map(str,c['source']['semesters']))})" for c in match['candidates']) or '—'
            lines.append(f"| {row['row']} | {clean(row['name'])} | {match['program']}-{match['intake']} | {match['status']} | {clean(sources)} |")
    lines += ['', '## Расхождения формул образца', '', '| Лист | Ячейка | Формула | Кэш | По сохранённым входам |', '|---|---|---|---:|---:|']
    for m in f['mismatches']:
        lines.append(f"| {clean(m['sheet'])} | {m['cell']} | {clean(m['formula'])} | {m['cached']} | {m['from_cached_inputs']} |")
    lines += ['', 'Циклические ячейки: ' + ', '.join(f'{s}!{c}' for s, c in f['cycles']) + '.', '',
              '## Контингент образца и выбранных групп', '',
              '| Строка образца | Предмет | Студенты в образце | В выбранных группах |', '|---:|---|---:|---:|']
    for row in result['template_rows']:
        if row['expected_students'] is not None and row['expected_students'] != row['students']:
            lines.append(f"| {row['row']} | {clean(row['name'])} | {row['students']} | {row['expected_students']} |")
    lines += ['', '## Проверка заданных правил СРСП и экзамена в образце', '',
              'Сравнение с действующими правилами проекта: СРСП = 1,5 × кредиты; экзамен = 0,1 × студенты. Пустая ячейка показана как None.', '',
              '| Ячейка образца | Сохранено | По правилу проекта |', '|---|---:|---:|']
    for row in result['template_rows']:
        for norm in row['norm_differences']:
            lines.append(f"| {norm['cell']} | {norm['actual']} | {norm['expected']} |")
    lines += ['', '## Сопоставление аудиторных часов', '',
              'Различие часов не всегда ошибка: в образце могут быть подгруппы и другой состав потока. Без подтверждённого сценария часы из образца не переносятся.', '',
              '| Строка образца | Исходный МУП:строка | Часы МУПа ЛК/ПЗ/ЛЗ | Часы образца | Сравнение |', '|---:|---|---|---|---|']
    for row in result['template_rows']:
        for match in row['matches']:
            for candidate in match['candidates']:
                src = candidate['source']
                status = 'Не определены по семестру' if candidate['hours_match'] is None else ('Совпадают' if candidate['hours_match'] else 'Различаются')
                lines.append(f"| {row['row']} | {clean(src['file'])}:{src['row']} | {src['hours']} | {row['hours'][:3]} | {status} |")
    lines += ['',
              'Полный перечень исходных строк, кредитов, семестров и часов сохранён в соседнем JSON-файле.']
    (directory / 'Проверка МУПов и образца.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return result


def write_kazakh_comparison(directory, template_rows, results, config):
    """Locate each Kazakh template subject in the actual generated workbooks."""
    locations = {}
    section_counts = {}
    for result in results:
        path = Path(result['file'])
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            term = None
            for row in wb['Расчет']:
                h = normal(row[0].value).replace('і', 'i')
                if h in ('iсеместр', 'iiсеместр'):
                    term = 1 if h == 'iсеместр' else 2
                if not row[1].comment or row[5].value not in ('KZ', 'RU', 'ENG', 'DE'):
                    continue
                if row[5].value == 'KZ':
                    key = (result['year'], term)
                    section_counts[key] = section_counts.get(key, 0) + 1
                for group, source_row, sem in re.findall(r'^([^:\n]+):.*?строка (\d+); семестр (\d+);', row[1].comment.text, re.M):
                    locations[group, int(source_row), int(sem)] = dict(
                        year=result['year'], file=path.name, row=row[0].row, language=row[5].value)
        finally:
            wb.close()
    selected = {(g['program'], g['intake']) for g in config['groups']}
    records = []
    for template in template_rows:
        if template['language'] != 'KZ':
            continue
        found, elsewhere = [], []
        for match in template['matches']:
            if (match['program'], match['intake']) not in selected:
                continue
            for candidate in match['candidates']:
                source = candidate['source']
                for sem in source['semesters']:
                    loc = locations.get((source['group'], source['row'], sem))
                    if loc:
                        (found if sem == template['semester'] and loc['year'] == 2026 else elsewhere).append(loc)
        found = list({(x['file'], x['row']): x for x in found}.values())
        elsewhere = list({(x['file'], x['row']): x for x in elsewhere}.values())
        if any(x['language'] == 'KZ' for x in found):
            status = 'Есть в казахском блоке'
        elif found:
            status = 'Находится в другом языковом блоке'
        elif not any((m['program'], m['intake']) in selected for m in template['matches']):
            status = ('Не выбраны группы набора 2025' if all(m['intake'] == 2025 for m in template['matches'])
                      else 'Не выбраны группы из шаблона')
        elif elsewhere:
            status = 'В МУПе другой семестр'
        else:
            status = 'Точного соответствия в выбранных МУПах нет'
        records.append(dict(template_row=template['row'], name=template['name'], term=template['term'],
                            status=status, locations=found or elsewhere,
                            groups=[f"{m['program']}-{m['intake']}" for m in template['matches']]))
    directory = Path(directory)
    payload = dict(groups=[{k: g[k] for k in ('name', 'program', 'intake', 'students')} for g in config['groups']], rows=records,
                   output_counts={f'{year}:{term}': count for (year, term), count in section_counts.items()})
    (directory / 'Казахский блок — сверка с шаблоном.json').write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    def clean(v):
        return ' '.join(str(v).split()).replace('|', '/')
    lines = ['# Какие предметы казахского блока отличаются от шаблона', '',
             'Сверка относится к книге 2026–2027. В шаблоне 24 предметные строки в I семестре и 30 во II. Число строк включает потоки разных курсов и наборов.', '',
             'Выбраны группы: ' + ', '.join(g['name'] for g in config['groups']) + '.', '',
             'Названия в новой книге взяты из МУПов без служебных подписей. Английский и немецкий языки находятся в иностранном блоке; остальные предметы и составляющие — в казахском или русском по настройке группы. Общий бюджет блока остаётся учтённым один раз; связи составляющих указаны в примечаниях.', '']
    for term in (1, 2):
        chosen = [r for r in records if r['term'] == term]
        counts = Counter(r['status'] for r in chosen)
        lines += [f'## {term} семестр', '',
                  f"В шаблоне {len(chosen)} строк; в новой книге {section_counts.get((2026, term), 0)} видимых строк казахского блока. Повторные строки одного предмета для разных групп и составляющие общих блоков входят в это число.", '',
                  '; '.join(f'{k}: {v}' for k, v in counts.items()) + '.', '',
                  '| Строка шаблона | Предмет в шаблоне | Почему есть или отсутствует | Место в расчёте |', '|---:|---|---|---|']
        for r in chosen:
            locs = '; '.join(f"{x['file']}:{x['row']} ({x['language']})" for x in r['locations']) or ', '.join(r['groups'])
            lines.append(f"| {r['template_row']} | {clean(r['name'])} | {r['status']} | {clean(locs)} |")
    (directory / 'Казахский блок — сверка с шаблоном.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return payload
