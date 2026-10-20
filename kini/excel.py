from __future__ import annotations

import ast
from copy import copy
import math
import operator
from pathlib import Path
import re
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties
from .model import normal, subject_language


class Rows:
    def __init__(self, wb, source, name):
        self.source = source
        index = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
        if name in wb.sheetnames:
            del wb[name]
        self.ws = wb.create_sheet(name, index)
        self.r = 0
        for c, d in source.column_dimensions.items():
            self.ws.column_dimensions[c] = copy(d)
        for prop in ['sheet_format', 'sheet_properties', 'page_setup', 'page_margins', 'print_options']:
            setattr(self.ws, prop, copy(getattr(source, prop)))
        self.ws.sheet_view.showGridLines = False

    def block(self, start, end=None):
        end = end or start
        offset = self.r + 1 - start
        for sr in range(start, end + 1):
            tr = sr + offset
            self.ws.row_dimensions[tr] = copy(self.source.row_dimensions[sr])
            self.ws.row_dimensions[tr].index = tr
            for c in range(1, min(self.source.max_column, 42) + 1):
                old = self.source.cell(sr, c)
                if isinstance(old, MergedCell):
                    continue
                new = self.ws.cell(tr, c)
                new._style = copy(old._style)
                new.alignment = copy(old.alignment)
                new.value = old.value if old.data_type != 'f' else None
                if old.comment:
                    new.comment = copy(old.comment)
        for m in self.source.merged_cells.ranges:
            if m.min_row >= start and m.max_row <= end:
                self.ws.merge_cells(start_row=m.min_row + offset, end_row=m.max_row + offset,
                                    start_column=m.min_col, end_column=m.max_col)
        self.r += end - start + 1
        return self.r if start == end else (start + offset, end + offset)

    def clear(self, row, first=1, last=20):
        for c in range(first, last + 1):
            cell = self.ws.cell(row, c)
            if not isinstance(cell, MergedCell):
                cell.value = None
                cell.comment = None

    def finish(self):
        self.ws.print_area = f'A1:T{self.r}'
        self.ws.page_setup.orientation = 'landscape'
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True


def sum_rows(ws, target, rows, first=8, last=20):
    for c in range(first, last + 1):
        ws.cell(target, c, '=' + '+'.join(f'{col(c)}{r}' for r in rows) if rows else '=0')


def totals(builder, template_row, rows, label=None):
    r = builder.block(template_row)
    if label is not None:
        builder.ws.cell(r, 1, label)
    sum_rows(builder.ws, r, rows)
    return r


def source_comment(load):
    lines = [f'{m.group}: {m.file} / {m.sheet}, строка {m.row}; семестр {m.semester}; '
             f'{m.credits} кр.; {m.students} студентов.' for m in load.members]
    lines.append('Часы: ' + load.hours_origin)
    for m in load.members:
        if m.options:
            lines.append('Варианты общего блока: ' + ' / '.join(m.options))
        lines.extend(m.notes)
    return '\n'.join(dict.fromkeys(lines))


def fit_subject_title(ws, row):
    """Keep long multilingual titles within Excel's maximum row height."""
    cell = ws.cell(row, 2)
    initial = int(cell.font.sz or 16)
    for size in range(initial, 7, -1):
        width = max(15, (ws.column_dimensions['B'].width or 30) * 10 / size)
        lines = sum(max(1, math.ceil(len(line) / width)) for line in str(cell.value).split('\n'))
        height = max(42, lines * size * 1.3 + 6)
        if height <= 409:
            font = copy(cell.font)
            font.sz = size
            cell.font = font
            ws.row_dimensions[row].height = height
            return
    raise ValueError(f'Название в строке {row} слишком длинное для печати в Excel; сократите текст названия.')


def fill_load(builder, load, sequence):
    r = builder.block(19)
    builder.clear(r)
    ws = builder.ws
    load.output_row = r
    name = load.members[0].name
    ws.cell(r, 1, sequence)
    ws.cell(r, 2, name)
    ws.cell(r, 3, load.credits)
    ws.cell(r, 4, load.members[0].course)
    ws.cell(r, 5, ', '.join(dict.fromkeys(m.group for m in load.members)))
    ws.cell(r, 6, load.language)
    ws.cell(r, 7, load.students)
    ws.cell(r, 2).comment = Comment(source_comment(load), 'KINI')
    for c in [2, 5]:
        ws.cell(r, c).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    fit_subject_title(ws, r)
    if load.kind == 'ordinary':
        for c, value in zip(range(8, 11), load.hours):
            ws.cell(r, c, value)
            if value is None:
                ws.cell(r, c).fill = PatternFill('solid', fgColor='FFF2CC')
                ws.cell(r, c).comment = Comment('Часы по этому семестру не определены. Подробности — в отчёте приложения.', 'KINI')
        ws.cell(r, 11, f'=ROUND(C{r}*1.5,6)')
        ws.cell(r, 12, f'=ROUND(G{r}*0.1,6)')
    elif load.kind == 'practice':
        ws.cell(r, 14, f'=ROUND(C{r}*G{r}*0.2,6)')
        ws.cell(r, 14).comment = Comment('Формула учебной практики из шаблона: Расчет!N25 и N59.', 'KINI')
    else:
        for c in [14, 16, 19]:
            ws.cell(r, c).fill = PatternFill('solid', fgColor='FFF2CC')
    # Явные настройки могут задать норму/часы; все участники объединённой строки
    # должны иметь одинаковые настройки (проверено на этапе группировки).
    for key, value in load.members[0].overrides.get('work_hours', {}).items():
        if key not in list('HIJKLMNOPQRS'):
            raise ValueError(f'Недопустимое поле work_hours: {key}')
        ws[f'{key}{r}'] = value
    ws.cell(r, 20, f'=SUM(H{r}:S{r})')
    for c in range(8, 21):
        ws.cell(r, c).number_format = '0.##'
    # Названия общего блока видны в книге, а не спрятаны в комментарии.
    # Пустые C/G/H:S этих строк означают общий бюджет в основной строке.
    return r


def fill_shared_subject(builder, load, title, sources, sequence):
    r = builder.block(19)
    builder.clear(r)
    ws = builder.ws
    ws.cell(r, 1, sequence)
    ws.cell(r, 2, title)
    ws.cell(r, 4, load.members[0].course)
    ws.cell(r, 5, ', '.join(dict.fromkeys(m.group for m, _ in sources)))
    ws.cell(r, 6, subject_language(title, load.language))
    ws.cell(r, 20, '=0')
    for column in (2, 5):
        ws.cell(r, column).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    fit_subject_title(ws, r)
    builder.shared_rows.append(r)
    builder.shared_links.append((r, load, sources))
    return r


def create_main(wb, source, loads, year, issues):
    b = Rows(wb, source, 'Расчет')
    b.shared_rows = []
    b.shared_links = []
    b.block(1, 12)
    ws = b.ws
    ws['N4'] = f'" _____ " _______________ {year} ж.'
    ws['A6'] = re.sub(r'202[0-9]\s*[-–]\s*202[0-9]', f'{year}–{year+1}', str(ws['A6'].value or ''))
    if f'{year}–{year+1}' not in str(ws['A6'].value):
        ws['A6'] = str(ws['A6'].value or '') + f'\nУчебный год {year}–{year+1}'
    ws.merge_cells('A7:T7')
    ws['A7'] = 'Предварительный расчёт'
    ws['A7'].font = Font(name='Arial', size=10, color='9C6500')
    ws['A7'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[7].height = 28
    language_totals, semester_totals = {}, {}
    detail_rows = {1: [], 2: []}
    for term in (1, 2):
        if term == 2:
            b.block(50, 52)
        b.block(13 if term == 1 else 53)
        counter = 0
        for language, heading, subtotal in [('KZ', 14, 39), ('RU', 40, 42), ('ENG', 43, 48)]:
            b.block(heading)
            rows = []
            entries = []
            for load in loads:
                if load.term != term:
                    continue
                block = 'ENG' if load.language in ('ENG', 'DE') else load.language
                if block == language:
                    entries.append((load.members[0].course, min(m.row for m in load.members),
                                    load.members[0].name, load, None))
                shared = {}
                for member in load.members:
                    for source_row in member.shared_sources:
                        shared.setdefault(source_row['name'], []).append((member, source_row))
                for title, sources in shared.items():
                    shared_language = subject_language(title, load.language)
                    block = 'ENG' if shared_language in ('ENG', 'DE') else shared_language
                    if block == language:
                        entries.append((load.members[0].course, min(s['row'] for _, s in sources), title, load, sources))
            for _, _, title, load, sources in sorted(entries, key=lambda e: e[:3]):
                counter += 1
                if sources is None:
                    rows.append(fill_load(b, load, counter))
                else:
                    fill_shared_subject(b, load, title, sources, counter)
            if not entries:
                empty = b.block(41)
                b.clear(empty)
                ws.cell(empty, 20, f'=SUM(H{empty}:S{empty})')
            detail_rows[term].extend(rows)
            language_totals[term, language] = totals(b, subtotal, rows)
        semester_totals[term] = totals(b, 49 if term == 1 else 95,
                                      [language_totals[term, lang] for lang in ('KZ', 'RU', 'ENG')])
    for row, load, sources in b.shared_links:
        lines = [f'{m.group}: {m.file} / {m.sheet}, строка {s["row"]}; семестр {s["semester"]}; '
                 f'{s["credits"]} кр. в исходной строке; общий бюджет — строка {load.output_row}.' for m, s in sources]
        ws.cell(row, 2).comment = Comment('\n'.join(lines), 'KINI')
    b.block(98)
    annual = totals(b, 99, list(semester_totals.values()))
    # Пустые разделы магистратуры сохраняются со всеми полями и корректными нулями.
    mag_start, mag_end = b.block(100, 124)
    for r in range(mag_start, mag_end + 1):
        original = 100 + r - mag_start
        for c in range(8, 21):
            if isinstance(ws.cell(r, c), MergedCell):
                continue
            if source.cell(original, c).data_type == 'f':
                ws.cell(r, c, '=0')
    b.block(125, 128)
    breakdown = {}
    for term, heading, start in [(1, 129, 130), (2, 134, 135), (0, 139, 140)]:
        b.block(heading)
        a, z = b.block(start, start + 3)
        for i, language in enumerate(('KZ', 'RU', 'ENG')):
            r = a + i
            for c in range(8, 21):
                cc = col(c)
                refs = ([language_totals[term, language]] if term else
                        [language_totals[t, language] for t in (1, 2)])
                ws.cell(r, c, '=' + '+'.join(f'{cc}{x}' for x in refs))
        sum_rows(ws, z, list(range(a, z)))
        breakdown[term] = z
    b.block(144, 147)
    allocation_rows = {}
    for key, title, source_title in [('main', None, 148), ('hourly', None, 152),
                                     ('unassigned', 'НЕ РАСПРЕДЕЛЕНО ПО ВИДУ ОПЛАТЫ / БӨЛІНБЕГЕН', 152)]:
        r = b.block(source_title)
        if title:
            ws.cell(r, 1, title)
        for term, sample in [(1, 149), (2, 150), (0, 151)]:
            r = b.block(sample)
            allocation_rows[key, term] = r
    b.block(156)
    allocated = {}
    for term, sample in [(1, 157), (2, 158), (0, 159)]:
        allocated[term] = totals(b, sample, [allocation_rows[k, term] for k in ('main', 'hourly', 'unassigned')])
    b.block(160, 164)
    b.block(167)
    checks = {}
    for term, label in [(1, '1 сем'), (2, '2 сем'), (0, 'Оқу жылы')]:
        r = b.block(169)
        b.clear(r)
        ws.cell(r, 7, label)
        for c in range(8, 21):
            ws.cell(r, c, f'={col(c)}{semester_totals[term] if term else annual}-{col(c)}{allocated[term]}')
        checks[term] = r
    b.finish()
    return dict(semesters=semester_totals, annual=annual, languages=language_totals,
                allocation=allocation_rows, allocated=allocated, checks=checks, detail=detail_rows,
                shared_detail=b.shared_rows)


def create_direction(wb, source, name, loads, year):
    b = Rows(wb, source, name)
    b.block(1, 12)
    ws = b.ws
    ws['A6'] = f'{name}\nУчебная нагрузка {year}–{year+1}'
    ws['N4'] = f'" _____ " _______________ {year} ж.'
    ws['A7'] = 'Данные связаны формулами с листом «Расчет». Незаданные часы отмечены в исходной строке.'
    ws.merge_cells('A7:T7')
    ws['A7'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[7].height = 26
    result = {}
    for term in (1, 2):
        b.block(13 if term == 1 else 53)
        rows = []
        for i, load in enumerate([x for x in loads if x.term == term], 1):
            r = b.block(19)
            b.clear(r)
            ws.cell(r, 1, i)
            for c in range(2, 21):
                ws.cell(r, c, f"='Расчет'!{col(c)}{load.output_row}")
            ws.cell(r, 2).comment = Comment(source_comment(load), 'KINI')
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[r].height = wb['Расчет'].row_dimensions[load.output_row].height
            rows.append(r)
        if not rows:
            r = b.block(41)
            b.clear(r)
        result[term] = totals(b, 49 if term == 1 else 95, rows)
    b.block(98)
    result[0] = totals(b, 99, [result[1], result[2]])
    b.block(160, 164)
    b.finish()
    return result


def quote_sheet(s):
    return "'" + s.replace("'", "''") + "'"


def create_summary(wb, source, name, directions, direction_totals):
    b = Rows(wb, source, name)
    b.block(1, 2)
    ws = b.ws
    rows_by_term = {1: [], 2: [], 0: []}
    for direction in directions:
        a, z = b.block(3, 5)
        ws.cell(a, 1, direction)
        for r, term in zip(range(a, z + 1), (1, 2, 0)):
            for c in range(2, 15):
                ws.cell(r, c, f'={quote_sheet(direction)}!{col(c+6)}{direction_totals[direction][term]}')
            rows_by_term[term].append(r)
            ws.cell(r, 2).comment = Comment({1: 'I семестр', 2: 'II семестр', 0: 'Учебный год'}[term], 'KINI')
    a, z = b.block(18, 20)
    ws.cell(a, 1, 'БАРЛЫҒЫ')
    result = {}
    for r, term in zip(range(a, z + 1), (1, 2, 0)):
        sum_rows(ws, r, rows_by_term[term], first=2, last=14)
        result[term] = r
    b.finish()
    ws.print_area = f'A1:N{b.r}'
    return result


def table_sheet(wb, name, headers, rows):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='365F91')
    for row in ws:
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[col(c)].width = 23 if c != len(headers) else 95
    return ws


def write_book(template, path, year, loads, main, hourly, issues):
    wb = openpyxl.load_workbook(template)
    source = openpyxl.load_workbook(template)
    structure = create_main(wb, source['Расчет'], loads, year, issues)
    # Итоговые листы временно отключены. Нижние блоки исходного Расчета
    # суммируют его собственные строки, без ссылок на удалённые листы.
    for payment in ('main', 'hourly', 'unassigned'):
        for term in (1, 2, 0):
            rows = [l.output_row for l in loads if l.payment == payment and (not term or l.term == term)]
            sum_rows(wb['Расчет'], structure['allocation'][payment, term], rows)
    for name in list(wb.sheetnames):
        if name != 'Расчет':
            del wb[name]
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True,
                                     calcMode='auto', iterate=False)
    wb.active = 0
    values = evaluate(wb)
    # Независимая сверка уровней агрегации перед сохранением.
    checks = validate(wb, loads, structure, values)
    wb.save(path)
    cache_formulas(path, wb, values)
    source.close()
    wb.close()
    serial_structure = {key: {str(k): v for k, v in value.items()} if isinstance(value, dict) else value
                        for key, value in structure.items()}
    return dict(year=year, file=str(path), load_rows=len(loads), contributions=sum(len(l.members) for l in loads),
                issues=sum(i.year in (None, year) for i in issues), totals=checks, structure=serial_structure,
                unassigned_rows=sum(l.payment == 'unassigned' for l in loads))


def create_direct_summary(wb, source, name, loads, directions):
    """Свод напрямую суммирует строки Расчета: промежуточные листы не нужны."""
    b = Rows(wb, source, name)
    b.block(1, 2)
    by_term = {1: [], 2: [], 0: []}
    for direction in dict.fromkeys([*directions, *sorted({l.direction for l in loads})]):
        original_row = next((r for r in range(3, source.max_row + 1)
                             if normal(source.cell(r, 1).value) == normal(direction)), None)
        sample_row = original_row or 3
        a, z = b.block(sample_row, sample_row + 2)
        b.ws.cell(a, 1, source.cell(original_row, 1).value if original_row else direction)
        for row, term in zip(range(a, z + 1), (1, 2, 0)):
            selected = [l for l in loads if l.direction == direction and (not term or l.term == term)]
            for c in range(2, 15):
                refs = [f"'Расчет'!{col(c+6)}{l.output_row}" for l in selected]
                b.ws.cell(row, c, '=' + '+'.join(refs) if refs else '=0')
            b.ws.cell(row, 2).comment = Comment({1: 'I семестр', 2: 'II семестр', 0: 'Учебный год'}[term], 'KINI')
            by_term[term].append(row)
    total_row = next(r for r in range(3, source.max_row + 1)
                     if 'барлы' in str(source.cell(r, 1).value or '').lower())
    last_direction = max(r for r in range(3, total_row) if source.cell(r, 1).value)
    if last_direction + 3 < total_row:
        b.block(last_direction + 3, total_row - 1)
    a, z = b.block(total_row, total_row + 2)
    b.ws.cell(a, 1, 'БАРЛЫҒЫ')
    result = {}
    for row, term in zip(range(a, z + 1), (1, 2, 0)):
        sum_rows(b.ws, row, by_term[term], first=2, last=14)
        result[term] = row
    if total_row + 3 <= source.max_row:
        b.block(total_row + 3, source.max_row)
    b.finish()
    b.ws.print_area = f'A1:N{b.r}'
    return result


REF = re.compile(r"(?:(?:'((?:[^']|'')+)'|([^!+*/(),\s]+))!)?(\$?[A-Z]{1,3}\$?\d+)")


def evaluate(wb):
    values, active = {}, set()
    ops = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}

    def arithmetic(node):
        if isinstance(node, ast.Constant):
            return node.value
        if isinstance(node, ast.BinOp):
            return ops[type(node.op)](arithmetic(node.left), arithmetic(node.right))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            return arithmetic(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == 'ROUND':
            return round(*(arithmetic(a) for a in node.args))
        raise ValueError(f'Неподдерживаемая формула: {ast.dump(node)}')

    def value(sheet, address):
        key = sheet, address.replace('$', '')
        if key in values:
            return values[key]
        if key in active:
            raise ValueError(f'Циклическая ссылка {sheet}!{address}')
        active.add(key)
        c = wb[sheet][address.replace('$', '')]
        result = c.value
        if c.data_type == 'f':
            expr = result[1:]
            m = REF.fullmatch(expr)
            if m:
                result = value((m[1] or m[2] or sheet).replace("''", "'"), m[3])
            else:
                def summed(m):
                    cells = wb[sheet][m[1]]
                    return str(sum(v for row in cells for cell in row
                                   if isinstance(v := value(sheet, cell.coordinate), (int, float))))
                expr = re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)', summed, expr)
                def referenced(m):
                    v = value((m[1] or m[2] or sheet).replace("''", "'"), m[3])
                    return repr(v if isinstance(v, (int, float)) else 0)
                expr = REF.sub(referenced, expr)
                result = arithmetic(ast.parse(expr, mode='eval').body)
        if result is None:
            result = 0
        values[key] = result
        active.remove(key)
        return result

    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    value(ws.title, cell.coordinate)
    return values


def validate(wb, loads, structure, values):
    ws = wb['Расчет']
    def val(r, c):
        cell = ws.cell(r, c)
        return values.get((ws.title, cell.coordinate), cell.value or 0)
    expected, exported = [], []
    for load in loads:
        for member in load.members:
            expected.append((member.group, member.row, member.semester))
            expected.extend((member.group, s['row'], s['semester']) for s in member.shared_sources)
    for row in ws:
        if row[1].comment:
            exported.extend((g, int(r), int(s)) for g, r, s in re.findall(
                r'^([^:\n]+):.*?строка (\d+); семестр (\d+);', row[1].comment.text, re.M))
    if len(expected) != len(set(expected)) or sorted(expected) != sorted(exported):
        raise AssertionError('Исходные дисциплины потеряны или повторены при записи Excel')
    for load in loads:
        r = load.output_row
        if val(r, 7) != sum({m.group: m.students for m in load.members}.values()):
            raise AssertionError('Контингент объединённого потока не совпадает')
        if load.kind == 'ordinary':
            if abs(val(r, 11) - load.credits * 1.5) > 1e-6:
                raise AssertionError('Неверный СРСП')
            if abs(val(r, 12) - load.students * .1) > 1e-6:
                raise AssertionError('Неверный экзамен')
        if abs(sum(val(r, c) for c in range(8, 20)) - val(r, 20)) > 1e-6:
            raise AssertionError('Итог строки не равен сумме видов работы')
    for c in range(8, 21):
        for term in (1, 2):
            actual = sum(val(l.output_row, c) for l in loads if l.term == term)
            if abs(actual - val(structure['semesters'][term], c)) > 1e-6:
                raise AssertionError('Итог семестра не равен сумме дисциплин')
        actual = sum(val(l.output_row, c) for l in loads)
        if abs(actual - val(structure['annual'], c)) > 1e-6:
            raise AssertionError('Итог года не равен сумме дисциплин')
        for r in structure['checks'].values():
            if abs(val(r, c)) > 1e-6:
                raise AssertionError('Распределённая нагрузка не сходится с общим расчётом')
    return {str(t): round(val(r, 20), 6) for t, r in {**structure['semesters'], 0: structure['annual']}.items()}


def cache_formulas(path, wb, values):
    """Пишет кэш только для формул, вычисленных выше; формулы сохраняются."""
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', ns)
    with zipfile.ZipFile(path) as archive:
        entries = {i.filename: (i, archive.read(i.filename)) for i in archive.infolist()}
    for i, ws in enumerate(wb, 1):
        filename = f'xl/worksheets/sheet{i}.xml'
        info, data = entries[filename]
        tree = ET.fromstring(data)
        for cell in tree.findall(f'.//{{{ns}}}c'):
            if cell.find(f'{{{ns}}}f') is None:
                continue
            v = cell.find(f'{{{ns}}}v')
            if v is None:
                v = ET.SubElement(cell, f'{{{ns}}}v')
            value = values[ws.title, cell.attrib['r']]
            if isinstance(value, str):
                cell.set('t', 'str')
                v.text = value
            else:
                cell.set('t', 'n')
                v.text = str(round(value, 10))
        entries[filename] = (info, ET.tostring(tree, encoding='utf-8'))
    temporary = Path(str(path) + '.tmp')
    with zipfile.ZipFile(temporary, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries.values():
            archive.writestr(info, data)
    temporary.replace(path)
