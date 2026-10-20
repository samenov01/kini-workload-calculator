from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
import re
import unicodedata

import openpyxl
from .layout import detect_layout, numeric, PlanValues


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def normal(value):
    value = unicodedata.normalize('NFKC', str(value or '')).lower().replace('ё', 'е')
    return re.sub(r'[^\w]+', '', value)


def tokens(name):
    return {normal(x) for x in re.split(r'[/|\n]', name) if normal(x)}


# Только варианты написания. Разные номера дисциплин не склеиваются.
ALIASES = {
    'history': ['история казахстана', 'history of kazakhstan'],
    'philosophy': ['философия', 'philosophy'],
    'sociology': ['социология', 'sociology'],
    'politics': ['политология', 'political science'],
    'culture': ['культурология', 'cultural studies'],
    'psychology': ['психология', 'psychology'],
    'math1': ['математика 1', 'mathematics 1', 'mathematik 1'],
    'math2': ['математика 2', 'mathematics 2', 'mathematik 2'],
    'math3': ['математика 3', 'mathematics 3', 'mathematik 3'],
    'statistics': ['статистика', 'statistics', 'математика 3 (статистика)',
                   'математика2(статистика)', 'mathematics 2 (statistics)',
                   'mathematics 2/3 (statistics)', 'mathematics 2 (сатистика)'],
    'kazakh': ['қазақ(орыс) тілі', 'казахский (русский) язык',
               'казахский язык (русский) язык', 'kazakh (russian) language'],
    'ict': ['информационные и коммуникационные технологии',
            'информационно-коммуникационные технологии',
            'информационно-коммуникационные технологии (е)',
            'информационно-коммуникационные технологии (на англ.языке)',
            'information and communication technologies',
            'information and communication technologies (in english)',
            'ақпараттық және коммуникациялық технологиялар',
            'ақпараттық және коммуникациялық технологиялар (е) (информатика 1)'],
    'german1': ['иностранный язык немецкий', 'foreign language german',
                'шет тілі неміс тілі а1', 'шет тілі неміс тілі a2'],
    'german2': ['иностранный язык немецкий 2 (а1-в1.2)',
                'foreign language german 2 (a1-b1.2)', 'шет тілі неміс тілі 2 (а1-в1. 2)'],
    'english': ['иностранный язык английский (а2.1-в1.2)',
                'foreign language english (a2.1-b1.2)',
                'шет тілі ағылшын тілі', 'шет тілі ағылшын тілі (а2.1-в1.2)'],
    'practice': ['учебная практика', 'educational practice', 'academic internship'],
    'physics1': ['физика 1', 'physics 1', 'physik 1'],
    'inorganic_chemistry': ['общая неорганическая химия', 'general inorganic chemistry'],
    'writing': ['написание научной работы и техника презентаций и риторика',
                'научная работа, техники презентации и риторика',
                'writing scientific papers and presentation techniques and rhetoric'],
    'chemistry_practicum': ['химия практикум', 'химия практикумы', 'chemistry practicum'],
    'economics_finance': ['основы экономики, предпринимательства и финансовой грамотности',
                          'основы экономики, препринимательства и финансовой грамотности',
                          'fundamentals of economics, entrepreneurship and financial literacy'],
}


def canonical(name, extra=None):
    parts = tokens(name)
    # Статистика проверяется до общего Mathematics 2.
    order = ['statistics'] + [k for k in ALIASES if k != 'statistics']
    for key in order:
        if parts & {normal(x) for x in ALIASES[key]}:
            return key
    for key, names in (extra or {}).items():
        if parts & {normal(x) for x in names}:
            return key
    return None


def subject_language(name, default='KZ'):
    """Foreign-language sections contain language subjects only.

    An English translation or a module heading does not set teaching language.
    """
    key = canonical(name)
    if key in ('german1', 'german2'):
        return 'DE'
    if key == 'english':
        return 'ENG'
    fallback = default if default in ('KZ', 'RU') else 'KZ'
    if key:
        return fallback
    for part in re.split(r'[/|\n]', name.lower()):
        part = part.strip()
        if re.match(r'(?:деловой немецкий|профессиональный (?:язык )?немецк|немецкий язык|'
                    r'шет тілі неміс|кәсіби (?:тіл )?неміс|іскери неміс|foreign language german|business german|professional language german)', part):
            return 'DE'
        if re.match(r'(?:иностранный язык английск|профессиональный (?:язык )?английск|английский язык|'
                    r'шет тілі ағылшын|кәсіби (?:тіл )?ағылшын|foreign language english|professional language english)', part):
            return 'ENG'
    return fallback


@dataclass
class Issue:
    kind: str
    group: str
    source: str
    row: int | str
    message: str
    year: int | None = None


@dataclass
class Course:
    group: str
    program: str
    intake: int
    students: int
    language: str
    file: str
    sheet: str
    row: int
    name: str
    code: str
    semester: int
    credits: float
    total_credits: float
    hours: tuple
    full_hours: tuple
    control: str
    kind: str
    key: str = ''
    options: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    overrides: dict = field(default_factory=dict)
    shared_sources: list = field(default_factory=list)
    hours_ambiguous: bool = False

    @property
    def year(self):
        return self.intake + (self.semester - 1) // 2

    @property
    def term(self):
        return 1 if self.semester % 2 else 2

    @property
    def course(self):
        return (self.semester + 1) // 2


def classify(name):
    n = name.lower()
    if any(t in n for t in ['учебная практика', 'оқу практикасы', 'academic internship']):
        return 'practice'
    if any(t in n for t in ['практика', 'практики', 'internship', 'практикa', 'дисертац',
                            'диплом', 'аттестаци', 'аттестац', 'нирс', 'ғылыми-зерттеу']):
        # Практикум по физике/химии не относится к руководству практикой.
        return 'special'
    return 'ordinary'


def read_plan(group, path, issues):
    wb = openpyxl.load_workbook(path, data_only=True)
    fw = openpyxl.load_workbook(path, data_only=False)
    if group.get('sheet'):
        ws = wb[group['sheet']]
    else:
        candidates = []
        for sheet in wb:
            try:
                detect_layout(sheet)
                candidates.append(sheet)
            except ValueError:
                pass
        if len(candidates) != 1:
            raise ValueError(f'{path.name}: найдено таблиц МУПа: {len(candidates)}. Выберите лист с корректными заголовками.')
        ws = candidates[0]
    fs = fw[ws.title]
    layout = detect_layout(ws)
    start, semcol = layout.header, layout.semesters[1]
    credit_col, controlcol = layout.credits, layout.control
    namecol, codecol = layout.name, layout.code
    resolver = PlanValues(wb, fw)
    merged = {}
    for m in ws.merged_cells.ranges:
        if m.max_row > m.min_row and m.max_row >= start:
            for rr in range(m.min_row, m.max_row + 1):
                for cc in range(m.min_col, m.max_col + 1):
                    merged[rr, cc] = (m.min_row, m.min_col)

    def val(r, c):
        address = ws.cell(*merged.get((r, c), (r, c))).coordinate
        return resolver.get(ws.title, address)

    def issue(kind, row, message):
        issues.append(Issue(kind, group['name'], path.name, row, message))

    courses = []
    module = ''
    language_route = False
    active_language_bases = {}
    language_parts = []
    by_anchor = defaultdict(list)
    excluded = set(group.get('exclude_rows', []))
    for r in range(start + 1, ws.max_row + 1):
        raw_name = ws.cell(r, namecol).value
        name = str(raw_name or '').strip()
        n = name.lower()
        is_module = any(t in n for t in ['модул', 'module', 'цикл ', 'компонент по выбору',
                                        'обязательный компонент', 'вузовский компонент',
                                        'дополнительные виды обучения']) or n.strip() in ('engineering', 'major')
        if (ws.cell(r, codecol).value and re.search(r'\d', str(ws.cell(r, codecol).value))
                and fs.cell(r, credit_col).data_type != 'f'):
            # EEE: одна дисциплина содержит слово «модулі» в переводе,
            # но имеет свой код, кредиты и семестр.
            is_module = False
        if is_module:
            module = name
            language_route = False
            active_language_bases.clear()
            continue
        if re.search(r'\bитого\b|барлығы|\btotal\b|келісілген|согласовано', n):
            module = ''
            language_route = False
            active_language_bases.clear()
            continue
        credits = numeric(val(r, credit_col))
        distribution = [numeric(val(r, c)) for c in layout.semesters.values()]
        present = [(sem, v) for sem, v in zip(layout.semesters, distribution) if number(v) and v > 0]
        if not present:
            if name and (ws.cell(r, codecol).value or number(credits)):
                # Альтернативы без собственных семестров сохраняются в журнале.
                if number(credits) or ws.cell(r, codecol).value:
                    issue('нет семестра', r, f'{name}: нет самостоятельного распределения по семестрам; в расчёт не добавлено.')
            continue
        if not number(credits):
            issue('нет кредитов', r, f'{name}: распределение есть, общее число кредитов не задано.')
            credits = sum(v for _, v in present)
        # Исключаем промежуточный итог без кода, когда рядом есть строка дисциплины.
        if not name and not ws.cell(r, codecol).value and fs.cell(r, credit_col).data_type == 'f':
            continue
        if not name:
            name = f'[Название отсутствует: {group["name"]}, строка {r}]'
            issue('нет названия', r, 'Кредиты и семестр сохранены; название дисциплины в МУПе пустое.')
        if r in excluded:
            issue('исключено настройкой', r, name)
            continue
        base_key = canonical(name)
        language_variant = (('german2' in active_language_bases and
                             any(t in n for t in ['деловой немецкий', 'профессиональный язык немецкий'])) or
                            ('english' in active_language_bases and
                             any(t in n for t in ['иностранный язык английский 2', 'публичных выступлений',
                                                  'профессиональный язык английский', 'business and soft skills',
                                                  'business  and soft skills'])))
        if language_variant:
            base = 'german2' if 'german2' in active_language_bases and any(
                t in n for t in ['деловой немецкий', 'профессиональный язык немецкий']) else 'english'
            language_parts.append((r, name, present, active_language_bases[base]))
            continue
        if base_key in ('german2', 'english'):
            active_language_bases[base_key] = r
        anchor = merged.get((r, credit_col), (r, credit_col))[0]
        if anchor != r and by_anchor[anchor]:
            if {s for s, _ in present} - {c.semester for c in by_anchor[anchor]}:
                raise ValueError(f'{path.name}, строка {r}: семестры общего блока не совпадают со строкой {anchor}.')
            for entry in by_anchor[anchor]:
                if entry.semester in dict(present):
                    entry.options.append(name)
                    entry.shared_sources.append(dict(row=r, name=name, code=str(ws.cell(r, codecol).value or ''),
                        semester=entry.semester, credits=dict(present)[entry.semester], kind='общие кредиты',
                        full_hours=[numeric(val(r, col)) for col in layout.hours]))
                    alternative_hours = entry.shared_sources[-1]['full_hours']
                    if (any(v is not None for v in alternative_hours)
                            and tuple(v or 0 for v in alternative_hours) != entry.full_hours
                            and 'hours' not in entry.overrides):
                        entry.hours = (None, None, None)
                        entry.hours_ambiguous = True
                        entry.notes.append('В составляющих общего блока указаны разные аудиторные часы; требуется уточнение вместо выбора первой строки.')
                        issue('часы общего блока', r, f'{name}: часы отличаются от общей строки {anchor}; аудиторные часы блока оставлены незаданными.')
            issue('общий блок МУПа', r, f'{name}: показан отдельной строкой внутри общего блока {anchor}; кредиты и нагрузка блока учитываются один раз.')
            continue
        full_hours = tuple(numeric(val(r, c)) for c in layout.hours)
        if any(number(x) for x in full_hours):
            full_hours = tuple(float(x) if number(x) else 0.0 for x in full_hours)
        else:
            full_hours = (None, None, None)
        if sum(v for _, v in present) != credits:
            issue('кредиты', r, f'{name}: всего {credits}, сумма по семестрам {sum(v for _, v in present)}. Используются семестровые значения.')
        kind = classify(name)
        for semester, ects in present:
            override = group.get('overrides', {}).get(f'{r}:{semester}', {})
            if override.get('exclude'):
                continue
            hours = full_hours if len(present) == 1 else (None, None, None)
            if 'hours' in override:
                hours = tuple(override['hours'])
            key = canonical(name)
            language = subject_language(override.get('name', name), override.get('language', group.get('language', 'KZ')))
            control = str(val(r, controlcol) or '')
            notes = []
            if len(present) > 1:
                notes.append('Часы МУПа заданы на несколько семестров; автоматическое пропорциональное деление не применяется.')
            c = Course(group['name'], group['program'], group['intake'], group['students'],
                       language, path.name, ws.title, r, override.get('name', name),
                       str(ws.cell(r, codecol).value or ''), semester, override.get('credits', ects),
                       credits, hours, full_hours, control, kind,
                       key=key or '', notes=notes, overrides=override)
            courses.append(c)
            by_anchor[anchor].append(c)
    # Языковой модуль может содержать общий маршрут и его поуровневые части.
    # Сохраняем каждую исходную строку и семестр; бюджет модуля не удваиваем.
    for r, name, present, base in language_parts:
        for semester, ects in present:
            targets = [c for c in courses if c.row == base and c.semester == semester]
            if not targets:
                raise ValueError(f'{path.name}, строка {r}: языковая часть имеет семестр {semester}, отсутствующий в общем маршруте {base}. Уточните структуру МУПа.')
            targets[0].shared_sources.append(dict(row=r, name=name, code=str(ws.cell(r, codecol).value or ''),
                semester=semester, credits=ects, kind='часть языкового блока',
                total_credits=numeric(val(r, credit_col)), semesters=len(present),
                full_hours=[numeric(val(r, col)) for col in layout.hours]))
        issue('общий языковой блок', r, f'{name}: сохранён в видимой части расчёта; общий бюджет строки {base} не начисляется повторно.')
        original_credits = numeric(val(r, credit_col))
        if original_credits != sum(v for _, v in present):
            issue('кредиты', r, f'{name}: в исходной строке {original_credits} кредитов, по семестрам {sum(v for _, v in present)}; расхождение сохранено в составе общего блока.')
    for c in courses:
        parts = [s for s in c.shared_sources if s['kind'] == 'часть языкового блока']
        if not parts:
            continue
        if sum(s['credits'] for s in parts) != c.credits:
            issue('структура языкового блока', c.row, f'{c.name}, семестр {c.semester}: части дают {sum(s["credits"] for s in parts)} кредитов, общий маршрут — {c.credits}; использован общий бюджет МУПа, требуется исправление источника.')
        # Переносим прямые односеместровые часы частей только если они
        # полностью согласованы с кредитами и общими часами маршрута.
        route = [s for entry in courses if entry.row == c.row for s in entry.shared_sources
                 if s['kind'] == 'часть языкового блока']
        valid = all(s['semesters'] == 1 and s['credits'] == s['total_credits']
                    and any(v is not None for v in s['full_hours']) for s in route)
        total_hours = tuple(sum(s['full_hours'][i] or 0 for s in route) for i in range(3))
        if (valid and total_hours == c.full_hours
                and all(sum(s['credits'] for s in entry.shared_sources if s['kind'] == 'часть языкового блока') == entry.credits
                        for entry in courses if entry.row == c.row)
                and 'hours' not in c.overrides):
            c.hours = tuple(sum(s['full_hours'][i] or 0 for s in parts) for i in range(3))
            c.notes.append('Часы семестра взяты из односеместровых частей языкового блока; сумма частей совпала с общими часами и кредитами МУПа.')
    # Итоги МУПа используются для диагностики, но не подменяют отдельные дисциплины.
    total_rows = [r for r in range(start + 1, ws.max_row + 1)
                  if re.search(r'\bитого\b|барлығы|\btotal\b', str(ws.cell(r, namecol).value or '').lower())
                  and number(ws.cell(r, credit_col).value)]
    if total_rows:
        r = total_rows[-1]
        computed = [sum(c.credits for c in courses if c.semester == sem) for sem in layout.semesters]
        stored = [ws.cell(r, c).value or 0 for c in layout.semesters.values()]
        if computed != stored:
            issue('итог МУПа', r, f'Суммы дисциплин по семестрам: {computed}; итоговая строка МУПа: {stored}. '
                  'Расчёт сформирован по дисциплинам, а не по ошибочным/ручным итогам.')
    wb.close()
    fw.close()
    if not courses:
        raise ValueError(f'{path.name}: не извлечено ни одной дисциплины.')
    return courses


def assign_keys(courses, extra):
    # Точное совпадение хотя бы одного перевода, без нечёткого поиска.
    alias_map = {}
    for c in courses:
        c.key = canonical(c.name, extra) or c.key
        if c.key:
            for t in tokens(c.name):
                alias_map.setdefault(t, set()).add(c.key)
    for c in courses:
        if not c.key:
            known = set().union(*(alias_map.get(t, set()) for t in tokens(c.name)))
            c.key = next(iter(known)) if len(known) == 1 else ''
        if not c.key:
            c.key = 'name:' + min(tokens(c.name))
        for t in tokens(c.name):
            alias_map.setdefault(t, set()).add(c.key)
    # Второй проход учитывает варианты, встретившиеся позже.
    for c in courses:
        known = set().union(*(alias_map.get(t, set()) for t in tokens(c.name)))
        if c.key.startswith('name:') and len(known) == 1:
            c.key = next(iter(known))
    return alias_map


def reference_key(name, courses, extra):
    key = canonical(name, extra)
    if key:
        return key
    matches = {c.key for c in courses if tokens(name) & tokens(c.name)}
    return next(iter(matches)) if len(matches) == 1 else None


def programs(value):
    # Кириллическое ЕЕЕ также встречается в шаблоне.
    value = str(value or '').upper().replace('Е', 'E')
    return set(re.findall(r'\b(CS|PE|LOG|MECH|EEE|IM)\s*-?\s*\d{2,4}\b', value))


def cohorts(value):
    value = str(value or '').upper().replace('Е', 'E')
    return {(p, int(y) + (2000 if len(y) == 2 else 0))
            for p, y in re.findall(r'\b(CS|PE|LOG|MECH|EEE|IM)\s*-?\s*(\d{2,4})\b', value)}


@dataclass
class Load:
    members: list[Course]
    hours: tuple
    hours_origin: str
    direction: str = 'Не распределено'
    payment: str = 'unassigned'
    output_row: int = 0
    combined_credits: float | None = None
    display_name: str = ''

    @property
    def credits(self): return self.combined_credits if self.combined_credits is not None else self.members[0].credits
    @property
    def students(self): return sum({x.group: x.students for x in self.members}.values())
    @property
    def term(self): return self.members[0].term
    @property
    def language(self): return self.members[0].language
    @property
    def year(self): return self.members[0].year
    @property
    def name(self): return self.display_name or self.members[0].name
    @property
    def kind(self): return self.members[0].kind


def build_loads(courses, template, config, issues):
    # Проверенные соответствия действуют только для конкретной программы,
    # названия и семестра. Они не меняют кредиты и период из МУПа.
    for c in courses:
        for rule in config.get('reference_matches', []):
            if (c.program == rule['program'] and c.semester == rule['semester']
                    and normal(c.name) == normal(rule['name'])):
                c.key = rule['key']
                c.notes.append('Соответствие образцу: ' + rule['evidence'])
    ref = openpyxl.load_workbook(template, data_only=True)
    ws = ref['Расчет']
    samples = []
    term = None
    for r in range(13, ws.max_row + 1):
        heading = normal(ws.cell(r, 1).value).replace('і', 'i')
        if heading in ('iсеместр', 'iiсеместр'):
            term = 1 if heading == 'iсеместр' else 2
        name = str(ws.cell(r, 2).value or '')
        if not term or not name or not number(ws.cell(r, 3).value) or not programs(ws.cell(r, 5).value):
            continue
        sample_cohorts = cohorts(ws.cell(r, 5).value)
        reference_counts = config.get('reference_students', {})
        reference_keys = {f'{p}-{year}' for p, year in sample_cohorts}
        valid_hours = True
        if reference_keys <= reference_counts.keys():
            expected = sum(reference_counts[k] for k in reference_keys)
            if ws.cell(r, 7).value != expected:
                issues.append(Issue('ошибка образца', '', Path(template).name, r,
                    f'Контингент образца {ws.cell(r, 7).value} не совпадает с подтверждённым составом {expected}; часы строки не используются.'))
                valid_hours = False
        samples.append(dict(key=reference_key(name, courses, config.get('aliases')),
                            term=term, course=ws.cell(r, 4).value,
                            credits=ws.cell(r, 3).value, students=ws.cell(r, 7).value,
                            language=ws.cell(r, 6).value, name=name,
                            programs=programs(ws.cell(r, 5).value), cohorts=cohorts(ws.cell(r, 5).value), row=r, valid_hours=valid_hours,
                            hours=tuple(ws.cell(r, i).value or 0 for i in (8, 9, 10))))
    # Направления и вид оплаты восстановлены из состава двух сводов шаблона.
    def directions(title):
        result = []
        for row in ref[title].iter_rows(min_row=3):
            label = row[0].value
            if 'барлы' in str(label or '').lower():
                break
            if not label:
                continue
            matches = [name for name in ref.sheetnames if normal(name) == normal(label)]
            if len(matches) != 1:
                raise ValueError(f'{title}!A{row[0].row}: направление «{label}» не найдено однозначно.')
            result.append(matches[0])
        return list(dict.fromkeys(result))
    main, hourly = directions('итог (основной)'), directions('итог (почас)')

    def eligible_samples(candidates, students):
        candidates = [s for s in candidates if s['valid_hours']]
        exact = [s for s in candidates if s['students'] == students]
        if exact:
            return exact
        # При росте контингента подтверждённые занятия не исчезают.
        # Это базовый объём; дополнительные подгруппы требуют отдельной нормы.
        return [s for s in candidates if number(s['students']) and 0 < s['students'] < students]

    def note_changed_students(sample, members, year):
        students = sum({m.group: m.students for m in members}.values())
        if sample['students'] != students:
            issues.append(Issue('численность потока', ', '.join(dict.fromkeys(m.group for m in members)),
                members[0].file, members[0].row,
                f'{members[0].name}: сохранён базовый объём занятий образца для {sample["students"]} студентов; '
                f'сейчас {students}. Дополнительные часы при делении на подгруппы требуют подтверждённой нормы.', year))
    if set(main) & set(hourly):
        raise ValueError('Направление одновременно включено в основной и почасовой своды шаблона.')
    routing = defaultdict(set)
    for sheet in main + hourly:
        s = ref[sheet]
        for row in s:
            if number(s.cell(row[0].row, 3).value) and isinstance(s.cell(row[0].row, 2).value, str):
                r = row[0].row
                key = reference_key(str(s.cell(r, 2).value), courses, config.get('aliases'))
                if key:
                    routing[key].add((sheet, 'main' if sheet in main else 'hourly'))
    # Во II семестре немецкий представлен двумя частями 4+2 кредита в МУПах,
    # а в образце — одним потоком на 6 кредитов (Расчет!90). Состав студентов
    # одинаковый: суммируем кредиты частей, но не повторяем людей и экзамен.
    consumed, loads = set(), []
    language_parts = defaultdict(list)
    for c in courses:
        if c.key in ('german1', 'german2') and c.language == 'DE' and not c.overrides:
            language_parts[c.year, c.term, c.course].append(c)
    for (year, term, course), parts in language_parts.items():
        grouped = defaultdict(list)
        for c in parts:
            grouped[c.group].append(c)
        if not grouped or not all(len(ms) == 2 and {c.key for c in ms} == {'german1', 'german2'} for ms in grouped.values()):
            continue
        credits = {sum(c.credits for c in ms) for ms in grouped.values()}
        if len(credits) != 1:
            continue
        credits = credits.pop()
        candidates = [s for s in samples if s['language'] == 'DE' and s['key'] in ('german1', 'german2')
                      and s['term'] == term and s['course'] == course and s['credits'] == credits
                      and s['cohorts'] == {(c.program, c.intake) for c in parts}]
        candidates = eligible_samples(candidates, sum(ms[0].students for ms in grouped.values()))
        if candidates and len({s['hours'] for s in candidates}) == 1:
            sample = candidates[0]
            route = routing.get(sample['key'], set())
            if len(route) != 1:
                continue
            direction, payment = next(iter(route))
            load = Load(parts, sample['hours'], f'Шаблон Расчет!H{sample["row"]}:J{sample["row"]}; '
                        'единый поток немецкого языка, части МУПа 4+2 кредита, студенты учтены один раз.',
                        direction, payment, combined_credits=credits, display_name=sample['name'])
            loads.append(load)
            note_changed_students(sample, parts, year)
            consumed.update(id(c) for c in parts)
            issues.append(Issue('части дисциплины', ', '.join(grouped), parts[0].file, parts[0].row,
                                f'Немецкий язык объединён в {credits} кредитов по строке {sample["row"]} образца; '
                                f'{sample["hours"][1]} практических часов, один экзамен на студента.', year))
    buckets = defaultdict(list)
    for c in courses:
        if id(c) not in consumed:
            buckets[c.year, c.term, c.course, c.key, c.credits, c.language, c.kind].append(c)
    for key, members in buckets.items():
        # Часы из шаблона допустимы только при том же сценарии потока.
        candidates = [s for s in samples if s['key'] == key[3] and s['term'] == key[1]
                      and s['course'] == key[2] and s['credits'] == key[4]
                      and s['language'] == key[5]
                      and s['cohorts'] == {(m.program, m.intake) for m in members}]
        candidates = eligible_samples(candidates, sum({m.group: m.students for m in members}.values()))
        if candidates and len({s['hours'] for s in candidates}) == 1 and not any(m.overrides or m.hours_ambiguous for m in members):
            origin = f'Шаблон Расчет!H{candidates[0]["row"]}:J{candidates[0]["row"]}; подтверждены дисциплина, курс, семестр, кредиты, язык и состав программ с годами набора.'
            note_changed_students(candidates[0], members, key[0])
            chunks = [(members, candidates[0]['hours'], origin)]
        else:
            by_hours = defaultdict(list)
            for m in members:
                # Одинаковые часы сами по себе не подтверждают общий поток.
                # Требуется строка образца с совместимыми наборами и периодом.
                confirmed_stream = any(s['key'] == m.key and s['term'] == m.term
                    and s['course'] == m.course and s['credits'] == m.credits
                    and s['language'] == m.language
                    and {(c.program, c.intake) for c in members} <= s['cohorts'] for s in samples)
                signature = (m.hours, m.full_hours if None in m.hours else (), m.total_credits if None in m.hours else 0,
                             tuple(sorted(m.overrides.get('work_hours', {}).items())),
                             tuple(m.options), '' if confirmed_stream else m.group)
                by_hours[signature].append(m)
            chunks = [(ms, sig[0], 'МУП: часы одной дисциплины в одном семестре' if None not in sig[0]
                       else 'Часы по семестру не определены') for sig, ms in by_hours.items()]
        for ms, hours, origin in chunks:
            if len({m.group for m in ms}) != len(ms):
                # Две записи одной группы никогда не схлопываем как разные группы.
                for m in ms:
                    loads.append(Load([m], m.hours, 'МУП: отдельная запись той же группы'))
                continue
            load = Load(ms, hours, origin)
            route = routing.get(ms[0].key, set())
            custom = config.get('routing', {}).get(ms[0].key)
            if custom:
                load.direction, load.payment = custom['sheet'], custom['payment']
            elif len(route) == 1:
                load.direction, load.payment = next(iter(route))
            # Учебная практика в шаблоне распределяется по программе.
            elif ms[0].kind == 'practice':
                if {m.program for m in ms} == {'LOG'}:
                    load.direction, load.payment = 'логистика', 'hourly'
                elif {m.program for m in ms} <= {'CS', 'MECH', 'PE', 'EEE'}:
                    load.direction, load.payment = 'comp science', 'main'
            if load.payment == 'unassigned':
                issues.append(Issue('распределение', ', '.join(m.group for m in ms), ms[0].file, ms[0].row,
                                    f'{load.name}: вид оплаты не определён; учтено в блоке «Не распределено» листа «Расчет».', load.year))
            if None in hours and load.kind == 'ordinary':
                issues.append(Issue('часы не определены', ', '.join(m.group for m in ms), ms[0].file, ms[0].row,
                                    f'{load.name}: лекции/практики/лабораторные требуют уточнения; СРСП и экзамен рассчитаны.', load.year))
            if load.kind == 'special':
                issues.append(Issue('норма не задана', ', '.join(m.group for m in ms), ms[0].file, ms[0].row,
                                    f'{load.name}: норма руководства/аттестации отсутствует; кредиты сохранены, преподавательские часы не выдумываются.', load.year))
            if len(ms) > 1 and not origin.startswith('Шаблон') and load.kind == 'ordinary':
                issues.append(Issue('общий поток', ', '.join(m.group for m in ms), ms[0].file, ms[0].row,
                                    f'{load.name}: состав потока подтверждён образцом для тех же наборов и периода; одинаковые часы МУПа учтены один раз. Нормы деления на подгруппы требуют проверки.', load.year))
            loads.append(load)
    ref.close()
    return loads, main, hourly
