from __future__ import annotations

from copy import copy, deepcopy
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from kini.excel import evaluate
from kini.audit import inventory_plan, verify_coverage, audit_formulas, compare_template
from kini.model import assign_keys, build_loads, read_plan, subject_language
from generate_kini import run


ROOT = Path(__file__).resolve().parent
FIXTURE_CONFIG = json.loads((ROOT / 'kini_config.json').read_text(encoding='utf-8'))
HAS_WORKBOOKS = all((ROOT / p).is_file() for p in
                    [FIXTURE_CONFIG['template'], *(g['file'] for g in FIXTURE_CONFIG['groups'])])


@unittest.skipUnless(HAS_WORKBOOKS, 'Нужны локальные МУПы и Excel-шаблон; см. README.')
class ParserAndMergeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = json.loads((ROOT / 'kini_config.json').read_text(encoding='utf-8'))
        cls.issues = []
        cls.courses = []
        for group in cls.config['groups']:
            cls.courses.extend(read_plan(group, ROOT / group['file'], cls.issues))
        assign_keys(cls.courses, {})
        cls.loads, _, _ = build_loads(cls.courses, ROOT / cls.config['template'], cls.config, cls.issues)

    def test_leaf_with_module_word_and_final_attestation_are_preserved(self):
        self.assertEqual(sum(c.credits for c in self.courses if c.group == 'EEE-26' and c.row == 92), 5)
        self.assertEqual(sum(c.credits for c in self.courses if c.group == 'PE-26' and c.row == 106), 8)
        self.assertFalse(any(c.group == 'LOG-26' and c.row == 99 for c in self.courses))
        self.assertEqual(sum(c.credits for c in self.courses if c.group == 'LOG-26' and c.semester == 8), 26)

    def test_blank_name_keeps_credits_and_semester(self):
        course = next(c for c in self.courses if c.group == 'LOG-26' and c.row == 78)
        self.assertEqual((course.credits, course.semester), (5, 6))
        self.assertIn('Название отсутствует', course.name)

    def test_shared_hours_do_not_turn_distinct_social_courses_into_options(self):
        courses = [c for c in self.courses if c.group == 'MECH-26' and c.row in (36, 37, 38, 39)]
        self.assertEqual(len(courses), 4)
        self.assertTrue(all(c.hours == (7, 15, 0) for c in courses))
        self.assertEqual(sum(c.credits for c in courses), 8)

    def test_shared_elective_block_is_counted_once(self):
        courses = [c for c in self.courses if c.group == 'MECH-26' and c.row in (51, 52, 53, 54)]
        self.assertEqual(len(courses), 1)
        self.assertEqual(len(courses[0].options), 3)
        self.assertEqual(courses[0].credits, 5)

    def test_language_routes_are_not_added_twice(self):
        self.assertFalse(any(c.group == 'PE-26' and c.row in (38, 39, 56, 57, 58, 59) for c in self.courses))
        self.assertEqual(sum(c.credits for c in self.courses if c.group == 'PE-26' and c.key == 'english'), 16)
        shared = {s['row'] for c in self.courses if c.group == 'PE-26' for s in c.shared_sources}
        self.assertTrue({38, 39, 56, 57, 58, 59} <= shared)

    def test_independent_source_inventory_includes_all_shared_subjects(self):
        inventory = []
        for group in self.config['groups']:
            inventory.extend(inventory_plan(ROOT / group['file'], group))
        coverage = verify_coverage(inventory, self.courses)
        self.assertEqual(coverage['missing'], [])
        self.assertEqual(coverage['extra'], [])
        self.assertEqual(coverage['duplicates'], [])
        self.assertEqual(coverage['expected'], 363)
        damaged = deepcopy(self.courses)
        next(c for c in damaged if c.shared_sources).shared_sources.pop()
        self.assertTrue(verify_coverage(inventory, damaged)['missing'])

    def test_equal_hours_without_a_confirmed_stream_remain_separate(self):
        a = deepcopy(self.courses[0])
        b = deepcopy(a)
        a.key = b.key = 'unknown_new_subject'
        a.group, b.group = 'A-26', 'B-26'
        a.hours = b.hours = (15, 30, 0)
        a.shared_sources = b.shared_sources = []
        loads, _, _ = build_loads([a, b], ROOT / self.config['template'], self.config, [])
        self.assertEqual(len(loads), 2)
        self.assertTrue(all(len(load.members) == 1 for load in loads))

    def test_template_hours_require_matching_language_and_intake(self):
        math = [deepcopy(c) for c in self.courses if c.key == 'math1']
        for changed_field, changed_value in [('language', 'RU'), ('intake', 2030)]:
            changed = deepcopy(math)
            for c in changed:
                setattr(c, changed_field, changed_value)
                c.hours = (1, 2, 3)
            loads, _, _ = build_loads(changed, ROOT / self.config['template'], self.config, [])
            self.assertTrue(all(load.hours == (1, 2, 3) for load in loads))
            self.assertFalse(any(load.hours_origin.startswith('Шаблон') for load in loads))

    def test_ambiguous_translation_is_not_assigned_arbitrarily(self):
        a, b, c = [deepcopy(self.courses[0]) for _ in range(3)]
        a.name, b.name, c.name = 'Alpha / Shared', 'Beta / Shared', 'Shared / Unique'
        a.key = b.key = c.key = ''
        assign_keys([a, b, c], {'first': ['Alpha'], 'second': ['Beta']})
        self.assertEqual((a.key, b.key), ('first', 'second'))
        self.assertNotIn(c.key, ('first', 'second'))

    def test_consistent_language_components_supply_direct_semester_hours(self):
        english = [c for c in self.courses if c.group == 'CS-26' and c.key == 'english']
        self.assertEqual([c.hours for c in english], [(30, 30, 0)] * 4)
        self.assertEqual(sum(c.credits for c in english), 16)

    def test_formula_audit_reports_circular_references(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'cycle.xlsx'
            wb = openpyxl.Workbook()
            wb.active['A1'] = '=SUM(A1:A2)'
            wb.active['A2'] = 3
            wb.active['B1'] = '=A2*2'
            title = wb.active.title
            wb.save(path)
            wb.close()
            audit = audit_formulas(path)
            self.assertIn((title, 'A1'), audit['cycles'])
            self.assertEqual(audit['unsupported'], [])
            self.assertEqual(audit['checked'], 2)

    def test_template_german_stream_matches_sum_of_source_parts(self):
        inventory = []
        for group in self.config['groups']:
            inventory.extend(inventory_plan(ROOT / group['file'], group))
        rows = compare_template(ROOT / self.config['template'], inventory, self.config)
        german = next(r for r in rows if r['row'] == 90)
        self.assertEqual(len(german['matches']), 5)
        self.assertTrue(all(m['status'] == 'Совпадает (сумма частей)' and m['combined_credits'] == 6
                            for m in german['matches']))

    def test_selecting_alternative_language_route_keeps_all_four_terms(self):
        group = deepcopy(next(g for g in self.config['groups'] if g['program'] == 'CS'))
        group['exclude_rows'] = [111]
        courses = read_plan(group, ROOT / group['file'], [])
        selected = [c for c in courses if c.row in (112, 113, 114, 115)]
        self.assertEqual({c.semester for c in selected}, {1, 2, 3, 4})
        self.assertEqual(sum(c.credits for c in selected), 16)
        self.assertEqual([c.language for c in selected], ['ENG', 'KZ', 'ENG', 'KZ'])

    def test_only_language_subjects_enter_foreign_language_section(self):
        cases = {
            'Информационно-коммуникационные технологии (на англ.языке)': 'KZ',
            'Навыки публичных выступлений (C1)': 'KZ',
            'Business and Soft Skills (B2.1-C1)': 'KZ',
            'Деловой немецкий язык (B2.2-C1.1)': 'DE',
            'Профессиональный язык английский (B2.1-C1)': 'ENG',
            'Физика / Physics': 'KZ',
        }
        for name, expected in cases.items():
            with self.subTest(name=name):
                self.assertEqual(subject_language(name, 'ENG'), expected)

    def test_reference_stream_sums_students_but_not_repeated_hours(self):
        math = [l for l in self.loads if l.year == 2026 and l.term == 1 and l.members[0].key == 'math1']
        self.assertEqual(len(math), 1)
        self.assertEqual(len(math[0].members), 5)
        self.assertEqual(math[0].students, 35)
        self.assertEqual(math[0].hours, (15, 60, 0))
        self.assertEqual(math[0].credits, 5)

    def test_annual_files_follow_intake_and_not_group_name(self):
        item = deepcopy(self.courses[0])
        item.intake = 2025
        item.semester = 4
        self.assertEqual((item.year, item.course, item.term), (2026, 2, 2))
        item.semester = 7
        self.assertEqual((item.year, item.course, item.term), (2028, 4, 1))

    def test_unknown_multisemester_hours_are_not_divided_arbitrarily(self):
        c = next(c for c in self.courses if c.group == 'MECH-26' and c.row == 43 and c.semester == 2)
        self.assertEqual(c.hours, (None, None, None))

    def test_every_contribution_occurs_in_exactly_one_load(self):
        parsed = {(c.group, c.row, c.semester) for c in self.courses}
        exported = [(c.group, c.row, c.semester) for l in self.loads for c in l.members]
        self.assertEqual(len(parsed), len(exported))
        self.assertEqual(parsed, set(exported))

    def test_german_parts_match_reference_without_doubling_students(self):
        german = [l for l in self.loads if l.year == 2026 and l.term == 2 and l.language == 'DE']
        self.assertEqual(len(german), 1)
        load = german[0]
        self.assertEqual((load.credits, load.students, load.hours), (6, 35, (0, 270, 0)))
        self.assertEqual(len(load.members), 10)
        for group in {c.group for c in load.members}:
            self.assertEqual(sorted(c.credits for c in load.members if c.group == group), [2, 4])

    def test_verified_chemistry_and_physics_streams(self):
        chemistry = [l for l in self.loads if l.year == 2026 and l.term == 1 and l.members[0].key == 'inorganic_chemistry']
        self.assertEqual(len(chemistry), 1)
        self.assertEqual({c.program for c in chemistry[0].members}, {'PE', 'EEE'})
        self.assertEqual((chemistry[0].students, chemistry[0].hours, chemistry[0].payment), (13, (15, 30, 30), 'hourly'))
        physics = [l for l in self.loads if l.year == 2026 and l.term == 2 and l.members[0].key == 'physics1']
        self.assertEqual(len(physics), 1)
        self.assertEqual({c.program for c in physics[0].members}, {'PE', 'MECH', 'EEE'})
        # Ошибочное число 208 в G61 образца не переносится в расчёт.
        self.assertEqual(physics[0].students, 20)
        self.assertEqual(physics[0].hours, (15, 15, 15))

    def test_reference_does_not_move_curriculum_semesters(self):
        self.assertEqual({c.semester for c in self.courses if c.group == 'LOG-26' and c.row == 67}, {4})
        self.assertEqual({c.semester for c in self.courses if c.group == 'CS-26' and c.row == 73}, {3})
        self.assertEqual({c.semester for c in self.courses if c.group == 'MECH-26' and c.row == 51}, {4})


@unittest.skipUnless(HAS_WORKBOOKS, 'Нужны локальные МУПы и Excel-шаблон; см. README.')
class GeneratedWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory()
        cls.manifest = run(ROOT / 'kini_config.json', output_dir=cls.directory.name)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def test_all_years_caches_styles_and_reconciliation(self):
        self.assertEqual([r['year'] for r in self.manifest['results']], [2026, 2027, 2028, 2029])
        config = json.loads((ROOT / 'kini_config.json').read_text(encoding='utf-8'))
        original = openpyxl.load_workbook(ROOT / config['template'])
        for result in self.manifest['results']:
            with self.subTest(year=result['year']):
                path = Path(result['file'])
                wb = openpyxl.load_workbook(path)
                cached = openpyxl.load_workbook(path, data_only=True)
                self.assertEqual(wb.sheetnames, ['Расчет'])
                self.assertEqual(wb['Расчет']['B10'].value, original['Расчет']['B10'].value)
                self.assertEqual(copy(wb['Расчет']['B10'].font), copy(original['Расчет']['B10'].font))
                self.assertEqual(copy(wb['Расчет']['B10'].border), copy(original['Расчет']['B10'].border))
                self.assertEqual(wb['Расчет'].column_dimensions['B'].width,
                                 original['Расчет'].column_dimensions['B'].width)
                values = evaluate(wb)
                for ws in wb:
                    for row in ws:
                        for cell in row:
                            if cell.data_type == 'f':
                                self.assertNotIn('!', cell.value, 'В однолистовой книге не должно быть межлистовых ссылок')
                                expected = values[ws.title, cell.coordinate]
                                actual = cached[ws.title][cell.coordinate].value
                                if isinstance(expected, (int, float)):
                                    self.assertAlmostEqual(actual, expected, places=6)
                                else:
                                    self.assertEqual(actual, expected)
                for r in result['structure']['checks'].values():
                    for c in range(8, 21):
                        self.assertAlmostEqual(cached['Расчет'].cell(r, c).value, 0, places=6)
                for term in (1, 2):
                    rows = result['structure']['detail'][str(term)]
                    total = result['structure']['semesters'][str(term)]
                    for c in range(8, 21):
                        expected = sum(cached['Расчет'].cell(r, c).value or 0 for r in rows)
                        self.assertAlmostEqual(cached['Расчет'].cell(total, c).value, expected, places=6)
                wb.close()
                cached.close()
        original.close()

    def test_every_shared_subject_is_visible_in_excel_and_not_counted_twice(self):
        expected = {(c['group'], s['row'], s['semester']): s['name']
                    for c in self.manifest['courses'] for s in c['shared_sources']}
        seen = []
        import re
        for result in self.manifest['results']:
            wb = openpyxl.load_workbook(result['file'], data_only=True)
            ws = wb['Расчет']
            for r in result['structure']['shared_detail']:
                cell = ws.cell(r, 2)
                self.assertIsNotNone(cell.comment)
                for g, source_row, sem in re.findall(r'^([^:\n]+):.*?строка (\d+); семестр (\d+);', cell.comment.text, re.M):
                    key = (g, int(source_row), int(sem))
                    self.assertEqual(expected[key], cell.value)
                    seen.append(key)
                self.assertTrue(all(ws.cell(r, col).value is None for col in (3, 7, *range(8, 20))))
                self.assertEqual(ws.cell(r, 20).value, 0)
            wb.close()
        self.assertEqual(set(seen), set(expected))
        self.assertEqual(len(seen), len(expected))

    def test_original_titles_and_language_sections(self):
        primary = {(c['group'], c['row'], c['semester']): c['name'] for c in self.manifest['courses']}
        for result in self.manifest['results']:
            wb = openpyxl.load_workbook(result['file'])
            ws = wb['Расчет']
            foreign_start = False
            for row in ws:
                heading = str(row[0].value or '')
                if 'тілінде оқыту' in heading and 'барлығы' not in heading:
                    foreign_start = heading.startswith('Ағылшын')
                if not row[1].comment or row[5].value not in ('KZ', 'RU', 'ENG', 'DE'):
                    continue
                title = row[1].value
                self.assertNotIn('Общий блок МУПа', title)
                self.assertNotIn('В составе блока', title)
                if foreign_start:
                    self.assertIn(subject_language(title), ('ENG', 'DE'))
                if row[0].row in result['structure']['detail'].get('1', []) + result['structure']['detail'].get('2', []):
                    import re
                    first = re.search(r'^([^:\n]+):.*?строка (\d+); семестр (\d+);', row[1].comment.text, re.M)
                    self.assertEqual(title, primary[first[1], int(first[2]), int(first[3])])
            wb.close()

    def test_student_change_flows_into_exam_semester_year_and_summary(self):
        result = self.manifest['results'][0]
        wb = openpyxl.load_workbook(result['file'])
        ws = wb['Расчет']
        row = next(r for r in result['structure']['detail']['1'] if ws.cell(r, 7).value == 35)
        before = evaluate(wb)
        ws.cell(row, 7, 45)
        after = evaluate(wb)
        self.assertAlmostEqual(after['Расчет', f'L{row}'] - before['Расчет', f'L{row}'], 1)
        for r in (result['structure']['semesters']['1'], result['structure']['annual']):
            self.assertAlmostEqual(after['Расчет', f'T{r}'] - before['Расчет', f'T{r}'], 1)
        for r in result['structure']['checks'].values():
            for c in range(8, 21):
                self.assertAlmostEqual(after['Расчет', ws.cell(r, c).coordinate], 0, places=6)
        wb.close()

    def test_larger_cohort_preserves_confirmed_language_hours(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            config = json.loads((ROOT / 'kini_config.json').read_text())
            config['template'] = str(ROOT / config['template'])
            counts = {'CS': 15, 'EEE': 11, 'LOG': 13, 'MECH': 13, 'PE': 12}
            for group in config['groups']:
                group['file'] = str(ROOT / group['file'])
                group['students'] = counts[group['program']]
            path = directory / 'config.json'
            path.write_text(json.dumps(config))
            manifest = run(path, output_dir=directory / 'output')
            result = manifest['results'][0]
            self.assertGreaterEqual(result['totals']['0'], self.manifest['results'][0]['totals']['0'])
            wb = openpyxl.load_workbook(result['file'], data_only=True)
            ws = wb['Расчет']
            rows = sum(result['structure']['detail'].values(), [])
            for language, expected in [('DE', 540), ('ENG', 360)]:
                hours = sum(ws.cell(r, col).value or 0 for r in rows
                            if ws.cell(r, 6).value == language for col in (8, 9, 10))
                self.assertEqual(hours, expected)
            wb.close()
            self.assertTrue(any(i['kind'] == 'численность потока' for i in manifest['issues']))

    def test_exams_with_missing_control_and_german_sums(self):
        result = self.manifest['results'][0]
        wb = openpyxl.load_workbook(result['file'], data_only=True)
        ws = wb['Расчет']
        german = next(r for r in result['structure']['detail']['2'] if ws.cell(r, 6).value == 'DE')
        self.assertEqual([ws.cell(german, c).value for c in (3, 7, 9, 11, 12, 20)], [6, 35, 270, 9, 3.5, 282.5])
        chemistry = next(r for r in result['structure']['detail']['2'] if 'Химия практикум' in ws.cell(r, 2).value)
        self.assertEqual(ws.cell(chemistry, 12).value, .9)
        wb.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
