from pathlib import Path
from collections import defaultdict
import json
import re
import openpyxl
from kini.model import number, normal, reference_key, Course

ROOT = Path(__file__).resolve().parent
OLD = Path.home() / 'Downloads/KINI_НАГРУЗКА_2026-2027.xlsx'
REFERENCE = Path.home() / 'Downloads/Telegram Desktop/расчет KINI 2026-27 (10).xlsx'


def dump(path, destination):
    w = openpyxl.load_workbook(path, data_only=True)
    f = openpyxl.load_workbook(path)
    lines = [str(path), str(w.sheetnames)]
    for sn in w.sheetnames:
        s = w[sn]
        lines.append('SHEET ' + sn)
        for row in s:
            cells = []
            for c in row[:20]:
                if c.value is not None or f[sn][c.coordinate].data_type == 'f':
                    formula = f[sn][c.coordinate].value
                    cells.append(f'{c.column_letter}:{str(c.value).replace(chr(10), " / ")}' +
                                 (f' [{formula}]' if f[sn][c.coordinate].data_type == 'f' else ''))
            if cells:
                lines.append(str(row[0].row) + ' ' + ' | '.join(cells))
    destination.write_text('\n'.join(lines), encoding='utf-8')
    w.close()
    f.close()


def details(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Расчет']
    term, rows, totals = 0, [], {}
    for r in range(13, ws.max_row + 1):
        heading = normal(ws.cell(r, 1).value).replace('і', 'i')
        if heading in ('iсеместр', 'iiсеместр'):
            term = 1 if heading == 'iсеместр' else 2
        if heading in ('iсеместрбойыншабарлығы', 'iiсеместрбойыншабарлығы') and term not in totals:
            totals[term] = ws.cell(r, 20).value
        if isinstance(ws.cell(r, 2).value, str) and number(ws.cell(r, 3).value) and isinstance(ws.cell(r, 5).value, str):
            comment = ws.cell(r, 2).comment
            sources = set(re.findall(r'^([^:\n]+):.*?строка (\d+); семестр (\d+);', comment.text if comment else '', re.M))
            rows.append(dict(row=r, term=term, name=ws.cell(r, 2).value, credits=ws.cell(r, 3).value,
                             groups=ws.cell(r, 5).value, students=ws.cell(r, 7).value,
                             hours=[ws.cell(r, c).value or 0 for c in range(8, 20)],
                             total=ws.cell(r, 20).value, sources=sources))
    wb.close()
    return rows, totals


def report():
    destination = ROOT / 'comparison'
    corrected = ROOT / 'output_corrected/KINI_НАГРУЗКА_2026-2027.xlsx'
    old, old_totals = details(OLD)
    ref, ref_totals = details(REFERENCE)
    new, new_totals = details(corrected)
    manifest = json.loads((ROOT / 'output_corrected/manifest.json').read_text(encoding='utf-8'))
    courses = [Course(**c) for c in manifest['courses']]
    def fmt(value):
        return f'{value:.2f}'.rstrip('0').rstrip('.').replace('.', ',')
    def title(value):
        return value.replace('\n', ' ').replace('|', '/')
    lines = ['# Сравнение двух расчётов и исправления', '',
        'Сравнены именно переданные файлы из Downloads, а не только рабочие копии. '
        'Исправлены генератор, правила сопоставления и выгрузка сводов. Исходные Excel не изменены.', '',
        '## Итоги', '',
        '| Вариант | I семестр | II семестр | За год |', '|---|---:|---:|---:|']
    for label, totals in [('Переданная выгрузка', old_totals), ('Образец: сохранённые T49/T95', ref_totals), ('Исправленная выгрузка пяти групп 2026', new_totals)]:
        lines.append(f'| {label} | {fmt(totals[1])} | {fmt(totals[2])} | {fmt(sum(totals.values()))} |')
    lines += ['', '**Образец и выгрузка имеют разный состав.** В образце также есть CS-25, MECH-25, IM-25, LOG-25 и EEE-25, '
              'включая потоки со смешанными наборами. В выбранных входных данных — только PE-26, CS-26, LOG-26, MECH-26 и EEE-26. '
              'Добавлять старшие группы с выдуманным МУПом или численностью нельзя. Их нужно выбрать отдельными МУПами в приложении.', '',
              '## Что исправлено в программе', '',
              '1. Экзамен обычной дисциплины теперь всегда равен 0,1 × студенты, как указал пользователь. '
              'Пустая форма контроля EEE больше не обнуляет экзамен. СРСП остаётся 1,5 × кредиты, в том числе для немецкого языка, '
              'хотя в образце его СРСП пропущен.',
              '2. Немецкий язык II семестра: части 4 и 2 кредита каждого МУПа сопоставлены с 6-кредитным потоком Расчет!90. '
              'Восстановлены 270 практических часов. Суммируются кредиты частей; 35 студентов и экзамен считаются один раз. '
              'Сопоставление требует совпадения состава программ, контингента, курса, семестра и суммы кредитов. '
              'Неподтверждённые другие составы не получают эти часы автоматически.',
              '3. PE: «Химия 1» из строки 70 сопоставлена с общей неорганической химией Расчет!38; '
              'поток PE + EEE = 13 студентов, часы 15/30/30. MECH: «Физика» из строки 71 сопоставлена с «Физика 1» Расчет!61; '
              'объединены PE + EEE + MECH = 20 студентов. Эти два соответствия ограничены программой, точным названием и семестром '
              'и записаны в kini_config.json → reference_matches с основанием.',
              '4. Варианты названия научного письма сопоставлены по точным названиям; одинаковые часы LOG и MECH учитываются '
              'один раз. «Химия практикумы» распознаётся как «Химия практикум» для почасового свода. '
              '«Коммуникация и презентация» CS не объявляется автоматически тем же предметом: это другой основной вариант блока выбора.',
              '5. Основные и почасовые направления читаются из соответствующих сводов образца, а не из зашитого перечня и порядка листов. '
              'Сохранены все 5 основных и 8 почасовых строк направлений, в исходном порядке, даже если в году у направления нулевые часы. '
              'Почасовой свод копирует собственный шаблон. Дополнительные листы не создаются.',
              '6. Считывание строк образца больше не ограничено диапазонами 15:38, 55:84 и т. п.; используются заголовки семестров '
              'и признаки строки дисциплины.', '',
              '## Изменения конкретных строк выгрузки', '',
              'Строки связаны по примечаниям с источниками (группа, строка МУПа, семестр), а не только по похожим названиям. '
              'При объединении приведена сумма прежних строк и итог новой.', '',
              '| Семестр | Строки до → после | Дисциплина | Было, ч | Стало, ч | Разница, ч |',
              '|---|---|---|---:|---:|---:|']
    # Connected components preserve many-to-one/one-to-many provenance during a merge.
    remaining = set(range(len(old)))
    changes = []
    while remaining:
        oi, ni = {min(remaining)}, set()
        while True:
            sources = set().union(*(old[i]['sources'] for i in oi), *(new[i]['sources'] for i in ni))
            next_old = {i for i, r in enumerate(old) if r['sources'] & sources} | oi
            next_new = {i for i, r in enumerate(new) if r['sources'] & sources}
            if (next_old, next_new) == (oi, ni):
                break
            oi, ni = next_old, next_new
        remaining -= oi
        before, after = sum(old[i]['total'] for i in oi), sum(new[i]['total'] for i in ni)
        if abs(before-after) > 1e-6:
            a, b = [old[i] for i in sorted(oi)], [new[i] for i in sorted(ni)]
            changes.append(dict(term=a[0]['term'], old_rows=[r['row'] for r in a], new_rows=[r['row'] for r in b], before=before, after=after))
            lines.append(f'| {a[0]["term"]} | '+ ', '.join(str(r['row']) for r in a) + ' → ' + ', '.join(str(r['row']) for r in b)
                         + f' | {title(b[0]["name"] if b else a[0]["name"])} | {fmt(before)} | {fmt(after)} | {fmt(after-before)} |')
    lines += ['', '## Почему отдельные строки образца отсутствуют в I курсе', '',
              '| Образец, Расчет | Где найдено в МУПе | Решение |', '|---|---|---|',
              '| 20: экономика и право, I семестр | PE: R50; MECH: R51 — IV семестр | Оставлено на II курс; не перенесено в I семестр ради совпадения |',
              '| 21: предпринимательство, I семестр | MECH: строка 53 — вариант общего блока F51:F54, R51:R54 | Не добавлено вторым предметом поверх тех же 5 кредитов |',
              '| 22: программирование 1, CS-26 | CS: R73 — III семестр | В расчёте 2027–2028 |',
              '| 23: основы экономики CS | CS: S58 — IV семестр | В расчёте 2027–2028 |',
              '| 25: практика LOG, I семестр | LOG: S67 — IV семестр; H67 содержит «1» | Используется сетка кредитов, не номер в графе контроля |',
              '| 58: научное письмо CS | CS: строки 60–61 — общий блок, основной вариант «Коммуникация и презентация» | Сохранён основной вариант; не заменён альтернативой без выбора |',
              '| 64: экологическая политика, все пять групп | EEE: Q81 — II семестр; у других групп нет такого распределения на I курс | Не добавлены отсутствующие студенты |',
              '| 84: математика 2, все пять групп | MECH: Q69 — III семестр, вариант со статистикой | В текущем II семестре только PE, CS, LOG, EEE |',
              '| 62–63: химия практикум и органическая химия, по 5 кредитов | EEE: Q63/Q64 — по 3 кредита | СРСП по 4,5 часа, кредиты МУПа сохранены |',
              '| Нет физкультуры | Есть во всех пяти МУПах в I и II семестрах | Сохранена, поскольку пользователь просил все дисциплины МУПа; нормы аудиторных часов не придуманы |', '',
              '## Ошибки самого образца, которые не перенесены', '',
              '- Расчет!T39 = SUM(T16:T37): пропущены история (строка 15) и химия (38). Аналогичные пропуски есть в H:K.',
              '- Расчет!T85 = SUM(T56:T83): пропущены политология (55) и математика 2 (84).',
              '- Расчет!J85 = SUM(J15:J83), M85 = SUM(M15:M83) и O85:S85 захватывают I семестр, его итоги и номера столбцов в строке 52. '
              'Отсюда искусственные часы 13, 15, 16, 17, 18, 19 в пустых видах нагрузки.',
              '- Расчет!G61 = 208 студентов для PE-26 + MECH-26 + EEE-26. В выбранных группах 4 + 7 + 9 = 20. '
              'Образец даёт 20,8 часа экзамена; исправленный файл — 2 часа. Часы H:J этой строки не перенесены: '
              'контингент образца некорректен; используются одинаковые часы 15/15/15 из трёх МУПов.',
              '- Расчет!L38 = 2 при G38 = 13; по заданной формуле экзамен должен быть 1,3 часа.',
              '- « химия»!G15:J15 ссылаются на Расчет!36 (ТК-сети), хотя название строки — неорганическая химия из строки 38. '
              'Своды исправленного файла ссылаются прямо на соответствующие новые строки Расчет.',
              '- В образце T169 = −100,5 и T170 = −63: верхние итоги не совпадают с суммой основного и почасового фондов.', '',
              '## Независимая сумма строк образца', '',
              'Здесь сложены только H:S строк дисциплин образца — без промежуточных итогов и шапок. '
              'Это диагностика образца, а не утверждённый расчёт: сами строки содержат описанные ошибки контингента и состава.', '',
              '| Семестр | Сумма H:S дисциплин образца | Сохранённый итог T49/T95 |', '|---|---:|---:|']
    for term in (1, 2):
        total = sum(sum(r['hours']) for r in ref if r['term'] == term)
        lines.append(f'| {term} | {fmt(total)} | {fmt(ref_totals[term])} |')
    lines += ['', '## Что остаётся неопределённым', '',
              'В исправленном расчёте по-прежнему явно отмечены данные, которых нет в источниках: '
              'часы PE по научному письму, распределение часов физкультуры по семестрам, '
              'нормы руководства дипломами/отдельными практиками, непроверенные назначения в фонды. '
              'Одинаковый числовой итог с образцом не является целью, если для него пришлось бы переносить дисциплины '
              'между семестрами, добавлять невыбранные группы или копировать ошибочные суммы.', '',
              'Проверены формулы СРСП и экзамена, уникальное покрытие всех исходных записей, '
              'суммы H:T по семестрам и году, своды и отсутствие повторного контингента в объединённых потоках. '
              'Во всех четырёх исправленных книгах контрольные разности равны нулю, листов ровно три.']
    (destination / 'Сравнение и исправления.md').write_text('\n'.join(lines), encoding='utf-8')
    (destination / 'changes.json').write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding='utf-8')
    dump(corrected, destination / 'generated_after.txt')
    return lines


if __name__ == '__main__':
    out = ROOT / 'comparison'
    out.mkdir(exist_ok=True)
    dump(OLD, out / 'generated_before.txt')
    dump(REFERENCE, out / 'reference.txt')
    if (ROOT / 'output_corrected/manifest.json').exists():
        report()
