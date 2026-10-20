"""Локальное приложение KINI: python webapp.py --open"""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from dataclasses import asdict
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import re
import secrets
import threading
from urllib.parse import quote, unquote, urlsplit
import webbrowser
import zipfile

import openpyxl

from generate_kini import run
from kini.layout import detect_layout
from kini.model import read_plan

ROOT = Path(__file__).resolve().parent
MAX_UPLOAD = 12 * 1024 * 1024


def inspect_file(path):
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
            raise ValueError('Распакованный Excel превышает 80 МБ.')
        if 'xl/workbook.xml' not in z.namelist():
            raise ValueError('Файл не является книгой Excel .xlsx.')
    wb = openpyxl.load_workbook(path, data_only=True)
    candidates, errors = [], []
    try:
        for ws in wb:
            try:
                if ws.max_row > 30000 or ws.max_column > 500:
                    raise ValueError('Таблица слишком большая: допускается до 30 000 строк и 500 столбцов.')
                layout = detect_layout(ws)
                text = ' '.join(str(c.value or '') for row in ws.iter_rows(max_row=layout.header) for c in row).lower()
                name_text = path.name.lower()
                program = ''
                patterns = [('PE', r'process engineering'), ('CS', r'computer science|\bcs\b'),
                            ('LOG', r'logistics|логистик'), ('MECH', r'mechatronic|мехатроник'),
                            ('EEE', r'energy and environmental|\beee\b'), ('IM', r'industrial management')]
                for p, pattern in patterns:
                    if re.search(pattern, text) or re.search(pattern, name_text):
                        program = p
                        break
                years = re.findall(r'(?:прием|год поступления|түскен жылы|admission|entry year|zulassungsjahr)\s*[:–-]?\s*(20\d{2})', text)
                intake = int(years[0]) if years else None
                group = dict(name='Предпросмотр', program=program or 'OTHER', intake=intake or 2026,
                             students=1, language='KZ', sheet=ws.title)
                issues = []
                courses = read_plan(group, path, issues)
                candidates.append(dict(sheet=ws.title, program=program, intake=intake,
                                       semesters=max(c.semester for c in courses), records=len(courses),
                                       disciplines=len({c.row for c in courses} | {s['row'] for c in courses for s in c.shared_sources}),
                                       warnings=len(issues), layout=asdict(layout)))
            except ValueError as exc:
                errors.append(dict(sheet=ws.title, message=str(exc)))
    finally:
        wb.close()
    if not candidates:
        raise ValueError('МУП не распознан. ' + ' '.join(x['message'] for x in errors))
    return dict(sheets=candidates, other_sheets=errors)


class Application:
    def __init__(self, root=ROOT):
        self.root = root
        self.runtime = root / '.kini-app'
        self.runtime.mkdir(exist_ok=True)
        self.uploads = self.runtime / 'uploads'
        self.uploads.mkdir(exist_ok=True)
        self.config = json.loads((root / 'kini_config.json').read_text(encoding='utf-8-sig'))
        self.token = secrets.token_urlsafe(32)
        self.lock = threading.RLock()
        self.executor = ThreadPoolExecutor(max_workers=1)
        self.files, self.jobs, self.defaults = {}, {}, []
        self._loaded = False

    def register(self, path, name=None):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()[:24]
        with self.lock:
            if digest in self.files:
                return self.files[digest]
        metadata = inspect_file(path)
        item = dict(id=digest, name=name or path.name, path=str(path.resolve()), **metadata)
        with self.lock:
            self.files[digest] = item
        return item

    def initialize(self):
        with self.lock:
            if self._loaded:
                return
            for group in self.config['groups']:
                item = self.register(self.root / group['file'])
                data = deepcopy(group)
                data.pop('file', None)
                data.update(file_id=item['id'], sheet=group.get('sheet', item['sheets'][0]['sheet']))
                self.defaults.append(data)
            for path in self.uploads.glob('*/*.xlsx'):
                try:
                    self.register(path)
                except (ValueError, zipfile.BadZipFile):
                    pass
            self._loaded = True

    def state(self):
        self.initialize()
        settings = self.runtime / 'selection.json'
        selected = json.loads(settings.read_text(encoding='utf-8')) if settings.exists() else self.defaults
        selected = [g for g in selected if g.get('file_id') in self.files]
        with self.lock:
            return dict(token=self.token, files=[{k:v for k,v in f.items() if k != 'path'} for f in self.files.values()],
                        groups=selected, jobs=[self.public_job(j) for j in self.jobs.values()],
                        selection_revision=hashlib.sha256(json.dumps(selected, sort_keys=True).encode()).hexdigest()[:16],
                        template=Path(self.config['template']).name)

    def upload(self, name, data):
        if not name.lower().endswith('.xlsx'):
            raise ValueError('Выберите файл .xlsx. Старый формат .xls нужно пересохранить в Excel.')
        safe = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', name)[:180]
        digest = hashlib.sha256(data).hexdigest()[:24]
        directory = self.uploads / digest
        directory.mkdir(exist_ok=True)
        path = directory / safe
        path.write_bytes(data)
        try:
            item = self.register(path, name=safe)
        except Exception as exc:
            path.unlink(missing_ok=True)
            raise ValueError(f'Не удалось прочитать «{safe}»: {exc}') from exc
        return {k:v for k,v in item.items() if k != 'path'}

    def validate_selection(self, groups):
        if not isinstance(groups, list) or not 1 <= len(groups) <= 80:
            raise ValueError('Добавьте от 1 до 80 групп.')
        result, names = [], set()
        for data in groups:
            if not isinstance(data, dict):
                raise ValueError('Некорректные настройки группы.')
            file = self.files.get(data.get('file_id'))
            if not file:
                raise ValueError('Выбранный файл не найден. Загрузите МУП заново.')
            sheet = data.get('sheet')
            if sheet not in {s['sheet'] for s in file['sheets']}:
                raise ValueError(f'{file["name"]}: выберите распознанный лист МУПа.')
            name = str(data.get('name', '')).strip()
            if not name or len(name) > 60 or name.casefold() in names:
                raise ValueError('Названия групп должны быть заполнены и не должны повторяться.')
            names.add(name.casefold())
            for field, low, high in [('intake', 2000, 2200), ('students', 1, 10000)]:
                value = data.get(field)
                if not isinstance(value, int) or isinstance(value, bool) or not low <= value <= high:
                    raise ValueError(f'{name}: некорректное поле {"год поступления" if field == "intake" else "студенты"}.')
            if data.get('program') not in ('CS', 'PE', 'EEE', 'LOG', 'MECH', 'IM', 'OTHER'):
                raise ValueError(f'{name}: выберите программу.')
            if data.get('language') not in ('KZ', 'RU', 'ENG', 'DE'):
                raise ValueError(f'{name}: выберите язык.')
            # Сохраняем продвинутые настройки существующих групп, но пути берём
            # только из внутреннего реестра, а не из запроса браузера.
            base = next((g for g in self.defaults if g['file_id'] == data['file_id'] and g['name'] == name), {})
            group = {key: deepcopy(base.get(key, default)) for key, default in
                     [('students_by_year', {}), ('exclude_rows', []), ('overrides', {})]}
            group.update({k:data[k] for k in ('file_id', 'sheet', 'program', 'intake', 'students', 'language')})
            group['name'] = name
            result.append(group)
        return result

    def start(self, groups):
        groups = self.validate_selection(groups)
        with self.lock:
            if any(j['status'] in ('queued', 'running') for j in self.jobs.values()):
                raise ValueError('Расчёт уже выполняется. Дождитесь его завершения.')
            job_id = secrets.token_hex(10)
            job = dict(id=job_id, status='queued', message='Подготовка файлов', results=[], issues=[])
            self.jobs[job_id] = job
            (self.runtime / 'selection.json').write_text(json.dumps(groups, ensure_ascii=False, indent=2), encoding='utf-8')
        self.executor.submit(self.worker, job_id, groups)
        return self.public_job(job)

    def worker(self, job_id, groups):
        directory = self.runtime / 'runs' / job_id
        directory.mkdir(parents=True, exist_ok=True)
        config = deepcopy(self.config)
        config['template'] = str((self.root / config['template']).resolve())
        config['output'] = str(directory / 'output')
        config['groups'] = []
        for g in groups:
            entry = deepcopy(g)
            entry['file'] = self.files[entry.pop('file_id')]['path']
            config['groups'].append(entry)
        config_path = directory / 'config.json'
        config_path.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding='utf-8')
        def progress(message):
            with self.lock:
                self.jobs[job_id].update(status='running', message=message)
        try:
            progress('Чтение МУПов')
            manifest = run(config_path, progress=progress)
            out = directory / 'output'
            allowed = {}
            results = []
            for result in manifest['results']:
                p = Path(result['file'])
                allowed[p.name] = p
                results.append({**result, 'file':p.name, 'url':f'/api/jobs/{job_id}/files/{quote(p.name)}'})
            # Отчёт приложения не содержит локальных абсолютных путей.
            report = dict(results=results, issues=manifest['issues'], courses=manifest['courses'],
                          coverage=manifest['coverage'], template_checks=manifest['template_checks'])
            report_path = out / 'Замечания.json'
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
            for audit_name in ('Проверка МУПов и образца.md', 'Проверка МУПов и образца.json',
                               'Казахский блок — сверка с шаблоном.md', 'Казахский блок — сверка с шаблоном.json'):
                allowed[audit_name] = out / audit_name
            archive = out / 'KINI_расчеты.zip'
            with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
                for p in allowed.values():
                    z.write(p, p.name)
                z.write(report_path, report_path.name)
            allowed[archive.name] = archive
            allowed[report_path.name] = report_path
            with self.lock:
                self.jobs[job_id].update(status='complete', message='Файлы сформированы. Итоги проверены.',
                    results=results, issues=manifest['issues'], allowed=allowed,
                    archive=f'/api/jobs/{job_id}/files/{quote(archive.name)}',
                    report=f'/api/jobs/{job_id}/files/{quote(report_path.name)}')
        except Exception as exc:
            with self.lock:
                self.jobs[job_id].update(status='error', message=str(exc))

    @staticmethod
    def public_job(job):
        return deepcopy({k:v for k,v in job.items() if k != 'allowed'})


class Handler(BaseHTTPRequestHandler):
    server_version = 'KINI/1.0'

    def respond(self, status, data, content_type='application/json; charset=utf-8', filename=None):
        if not isinstance(data, bytes):
            data = json.dumps(data, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(data)))
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Content-Security-Policy', "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'")
        if filename:
            self.send_header('Content-Disposition', "attachment; filename*=UTF-8''" + quote(filename))
        self.end_headers()
        self.wfile.write(data)

    def allowed_host(self):
        port = self.server.server_address[1]
        return self.headers.get('Host') in (f'127.0.0.1:{port}', f'localhost:{port}')

    def do_GET(self):
        if not self.allowed_host():
            return self.respond(403, {'error':'Недопустимый адрес сервера.'})
        app = self.server.app
        path = unquote(urlsplit(self.path).path)
        try:
            if path == '/api/state':
                return self.respond(200, app.state())
            match = re.fullmatch(r'/api/jobs/([a-f0-9]{20})(?:/files/([^/]+))?', path)
            if match:
                with app.lock:
                    job = app.jobs.get(match[1])
                    if not job:
                        return self.respond(404, {'error':'Расчёт не найден. После перезапуска сервера сформируйте его заново.'})
                    if match[2]:
                        file = job.get('allowed', {}).get(match[2])
                        if not file:
                            return self.respond(404, {'error':'Файл не найден.'})
                        mime = 'application/zip' if file.suffix == '.zip' else 'application/octet-stream'
                        return self.respond(200, file.read_bytes(), mime, file.name)
                    return self.respond(200, app.public_job(job))
            files = {'/':'index.html', '/app.js':'app.js', '/style.css':'style.css'}
            if path in files:
                mime = {'/':'text/html', '/app.js':'text/javascript', '/style.css':'text/css'}[path]
                return self.respond(200, (app.root / 'web' / files[path]).read_bytes(), mime + '; charset=utf-8')
            return self.respond(404, {'error':'Страница не найдена.'})
        except (ValueError, OSError, zipfile.BadZipFile) as exc:
            return self.respond(400, {'error':str(exc)})

    def do_POST(self):
        app = self.server.app
        if not self.allowed_host() or not secrets.compare_digest(self.headers.get('X-KINI-Token', ''), app.token):
            return self.respond(403, {'error':'Обновите страницу и повторите действие.'})
        try:
            length = int(self.headers.get('Content-Length', '0'))
            if not 0 < length <= MAX_UPLOAD:
                return self.respond(413, {'error':'Размер запроса должен быть от 1 байта до 12 МБ.'})
            data = self.rfile.read(length)
            if self.path == '/api/upload':
                return self.respond(200, app.upload(unquote(self.headers.get('X-Filename', '')), data))
            if self.path == '/api/generate':
                payload = json.loads(data)
                if not isinstance(payload, dict):
                    raise ValueError('Некорректные параметры расчёта.')
                return self.respond(202, app.start(payload.get('groups')))
            return self.respond(404, {'error':'Действие не найдено.'})
        except (ValueError, KeyError, TypeError, OSError, zipfile.BadZipFile) as exc:
            return self.respond(400, {'error':str(exc)})

    def log_message(self, fmt, *args):
        pass


def serve(port=8765, open_browser=False):
    app = Application()
    server = ThreadingHTTPServer(('127.0.0.1', port), Handler)
    server.app = app
    address = f'http://127.0.0.1:{server.server_address[1]}'
    print('KINI: ' + address, flush=True)
    if open_browser:
        threading.Timer(.4, lambda: webbrowser.open(address)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
        app.executor.shutdown(wait=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--port', type=int, default=8765)
    parser.add_argument('--open', action='store_true')
    args = parser.parse_args()
    try:
        serve(args.port, args.open)
    except OSError as exc:
        print(f'Не удалось запустить KINI: {exc}. Попробуйте python webapp.py --open --port 8766')
        raise SystemExit(1)
