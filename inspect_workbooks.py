from pathlib import Path
from collections import Counter
import argparse, json, warnings
import openpyxl

warnings.simplefilter('ignore', UserWarning)
parser = argparse.ArgumentParser(description='Inspect local Excel workbooks.')
parser.add_argument('--source', type=Path, default=Path('input'))
source = parser.parse_args().source
out = Path('workbook_review')
out.mkdir(exist_ok=True)
summary = []
for path in sorted(source.glob('*.xlsx')):
    wb = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    book = {'file': path.name, 'sheets': [], 'external_links': len(wb._external_links),
            'defined_names': {k: str(v) for k,v in wb.defined_names.items()}}
    lines = []
    for ws in wb:
        rows = []
        formulas, errors, samples = [], [], []
        for row in ws:
            populated = []
            for cell in row:
                if cell.value is None: continue
                value = cell.value
                if cell.data_type == 'f':
                    cv = cached[ws.title][cell.coordinate].value
                    formulas.append({'cell':cell.coordinate, 'formula':value, 'cached':cv})
                    value = f'{value} [cached={cv}]'
                if cell.data_type == 'e' or cached[ws.title][cell.coordinate].data_type == 'e':
                    errors.append({'cell':cell.coordinate,'value':str(cached[ws.title][cell.coordinate].value)})
                populated.append(f'{cell.column_letter}: {value}')
            if populated:
                rows.append(f'{row[0].row}\t' + ' | '.join(populated))
        info = {'name':ws.title,'state':ws.sheet_state,'max_row':ws.max_row,'max_column':ws.max_column,
                'populated_rows':len(rows),'formula_count':len(formulas),'errors':errors,
                'merged':list(map(str,ws.merged_cells.ranges)),
                'hidden_rows':[i for i,d in ws.row_dimensions.items() if d.hidden],
                'hidden_columns':[i for i,d in ws.column_dimensions.items() if d.hidden],
                'tables':list(ws.tables),'validations':[str(v) for v in ws.data_validations.dataValidation],
                'formulas':formulas}
        book['sheets'].append(info)
        lines.extend([f'\n### SHEET: {ws.title}\n']+rows)
    (out/(path.stem+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
    (out/(path.stem+'.json')).write_text(json.dumps(book,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    summary.append(book)
(out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
for b in summary:
    print(b['file'], 'external_links=', b['external_links'])
    for s in b['sheets']:
        print(' ',s['name'],f"{s['max_row']}x{s['max_column']}", 'rows=',s['populated_rows'], 'formulas=',s['formula_count'],'errors=',len(s['errors']),s['state'])
